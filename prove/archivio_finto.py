"""
Documenti di casa **finti** per le prove dell'archivio (calliope/archivio/): bollette,
fatture, garanzie, polizze, referti, documenti d'identità e contratti di persone, aziende e
numeri inventati. Mai documenti veri nel repository.

`genera(cartella)` scrive i file (PDF con testo, PDF scansionati, foto JPG storte, Word) e
restituisce, per ognuno, il percorso, il testo e i valori attesi (la «verità» con cui le prove
misurano OCR ed estrazione). Si usa da prove/prova_archivio.py (a secco) e dalle misure sulla
DGX (prove/archivio_misura.py).

    python prove\\archivio_finto.py <cartella>     # solo per guardarli
"""

import io
import json
import random
import shutil
import sys
from pathlib import Path

# Ogni documento: nome del file, formato («pdf» con testo, «scansione» = PDF fatto di
# un'immagine, «foto» = JPG storto e rumoroso, «docx»), righe di testo, valori attesi
DOCUMENTI = [
    {"file": "bolletta luce settembre 2026.pdf", "formato": "pdf", "tipo": "bolletta",
     "righe": ["LUMEN ENERGIA S.p.A.", "Via Fittizia 1, 00100 Roma - P.IVA 01234567897",
               "BOLLETTA ENERGIA ELETTRICA n. 2026/0912345",
               "Intestatario: Mario Bianchi  Codice fiscale BNCMRA80A01F205X",
               "Fornitura: Via dei Tigli 12, 20100 Milano",
               "POD IT001E12345678",
               "Periodo di fatturazione: dal 01/08/2026 al 31/08/2026",
               "Data emissione: 10/09/2026",
               "Consumo fatturato: 210 kWh",
               "TOTALE DA PAGARE: 84,50 euro",
               "Scadenza pagamento: 15/10/2026",
               "Pagamento con domiciliazione bancaria"],
     "atteso": {"categoria": "luce", "emittente": "Lumen Energia", "totale": 84.50,
                "data": "2026-09-10", "scadenza": "2026-10-15", "intestatario": "Mario Bianchi"}},
    {"file": "luce 2025-09.pdf", "formato": "pdf", "tipo": "bolletta",
     "righe": ["Lumen Energia SpA", "Bolletta luce numero 2025/0954321",
               "Cliente: BIANCHI MARIO - C.F. BNCMRA80A01F205X",
               "Indirizzo di fornitura: Via dei Tigli 12, Milano",
               "Periodo: agosto 2025 (01/08/2025 - 31/08/2025)",
               "Emessa il 09/09/2025",
               "Consumo: 240 kWh",
               "Importo totale: € 97,20",
               "Da pagare entro il 15/10/2025"],
     "atteso": {"categoria": "luce", "emittente": "Lumen Energia", "totale": 97.20,
                "data": "2025-09-09", "scadenza": "2025-10-15", "intestatario": "Mario Bianchi"}},
    {"file": "gas ottobre.jpg", "formato": "foto", "tipo": "bolletta",
     "righe": ["GASNOVA S.R.L.", "Bolletta gas naturale n. G-778812",
               "Intestatario: Mario Bianchi", "Via dei Tigli 12 - Milano",
               "PDR 01234567890123", "Periodo: 01/08/2026 - 30/09/2026",
               "Data documento: 01/10/2026", "Consumo: 45 Smc",
               "Totale bolletta: 62,30 €", "Scadenza: 20/10/2026"],
     "atteso": {"categoria": "gas", "emittente": "Gasnova", "totale": 62.30,
                "data": "2026-10-01", "scadenza": "2026-10-20", "intestatario": "Mario Bianchi"}},
    {"file": "scontrino lavatrice.pdf", "formato": "scansione", "tipo": "ricevuta",
     "righe": ["ELETTRODOMESTICI ROSSI s.n.c.", "Corso Immaginario 45, Milano",
               "P.IVA 09876543217", "FATTURA n. 118/2025 del 12/03/2025",
               "Cliente: Mario Bianchi", "Lavatrice Candor LV800  1 x 449,00",
               "Trasporto e installazione  1 x 30,00",
               "TOTALE EURO 479,00", "Pagato con carta"],
     "atteso": {"emittente": "Elettrodomestici Rossi", "totale": 479.00, "data": "2025-03-12"}},
    {"file": "garanzia lavatrice.docx", "formato": "docx", "tipo": "garanzia",
     "righe": ["CERTIFICATO DI GARANZIA", "Prodotto: lavatrice Candor LV800",
               "Numero di serie: LV800-55A21", "Venditore: Elettrodomestici Rossi snc",
               "Acquirente: Mario Bianchi", "Data di acquisto: 12/03/2025",
               "Durata della garanzia: 24 mesi (fino al 12/03/2027)",
               "Conservare lo scontrino insieme a questo certificato."],
     "atteso": {"scadenza": "2027-03-12", "bene": "lavatrice Candor LV800"}},
    {"file": "polizza auto.pdf", "formato": "scansione", "tipo": "assicurazione",
     "righe": ["AURORA ASSICURAZIONI S.p.A.", "Polizza RC Auto n. AU-2026-445566",
               "Contraente: Mario Bianchi", "Veicolo: Fiat Panda targa AB123CD",
               "Decorrenza: 01/11/2025", "Scadenza: 01/11/2026",
               "Premio annuo: 520,00 euro", "Data di emissione: 20/10/2025"],
     "atteso": {"emittente": "Aurora Assicurazioni", "totale": 520.00,
                "scadenza": "2026-11-01", "numero": "AU-2026-445566"}},
    {"file": "referto Giulia.jpg", "formato": "foto", "tipo": "referto",
     "righe": ["POLIAMBULATORIO SAN LUCA", "Paziente: Giulia Bianchi",
               "Data di nascita: 14/07/2012", "Esame: emocromo completo",
               "Data esame: 05/09/2026", "Emoglobina 13,2 g/dL (12,0 - 16,0)",
               "Globuli bianchi 6,8 x10^3/uL (4,0 - 10,0)", "Medico: dott.ssa Anna Verdi"],
     "atteso": {"persona": "Giulia Bianchi", "data": "2026-09-05"}},
    {"file": "carta identita Giulia.jpg", "formato": "foto", "tipo": "identita",
     "righe": ["REPUBBLICA ITALIANA", "CARTA DI IDENTITA' ELETTRONICA",
               "Numero CA12345AB", "Cognome: BIANCHI  Nome: GIULIA",
               "Nata il 14/07/2012", "Comune di Milano",
               "Emissione: 14/07/2021", "Scadenza: 14/07/2031"],
     "atteso": {"persona": "Giulia Bianchi", "scadenza": "2031-07-14", "numero": "CA12345AB"}},
    {"file": "contratto internet.pdf", "formato": "pdf", "tipo": "contratto",
     "righe": ["RETE VELOCE S.p.A.", "Contratto di abbonamento fibra n. RV-99812",
               "Cliente: Mario Bianchi, Via dei Tigli 12, Milano",
               "Canone mensile: 29,90 euro", "Data di attivazione: 01/02/2026",
               "Durata: 24 mesi, con rinnovo tacito",
               "Disdetta con preavviso di 30 giorni", "Firmato il 20/01/2026"],
     "atteso": {"emittente": "Rete Veloce", "totale": 29.90, "data": "2026-01-20"}},
    # Impaginata a due colonne, testo piccolo: l'ordine di lettura conta
    {"file": "acqua settembre.jpg", "formato": "foto", "tipo": "bolletta",
     "righe": ["ACQUE DEL NAVIGLIO S.p.A.",
               "Fattura servizio idrico n. AN-55120 | Data: 05/09/2026",
               "Intestataria: Laura Bianchi | Codice cliente: 7781203",
               "Voce | Importo (euro)",
               "Quota fissa | 12,40",
               "Consumi 38 mc | 41,85",
               "Fognatura e depurazione | 18,30",
               "IVA 10% | 7,26",
               "Totale fattura | 79,81",
               "Scadenza: 30/09/2026"],
     "atteso": {"categoria": "acqua", "emittente": "Acque del Naviglio", "totale": 79.81,
                "data": "2026-09-05", "scadenza": "2026-09-30", "intestatario": "Laura Bianchi"}},
]


def testo(doc: dict) -> str:
    return "\n".join(doc["righe"])


def _font(size: int):
    from PIL import ImageFont
    for name in ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def immagine(doc: dict, storta: bool = False, seme: int = 1):
    """Il documento come immagine: pagina bianca a 150 dpi circa; `storta` = foto fatta
    a mano (rotazione, rumore, luce non uniforme)."""
    from PIL import Image, ImageDraw, ImageFilter
    rnd = random.Random(seme)
    w, h = 1240, 140 + 56 * len(doc["righe"]) + 120
    img = Image.new("L", (w, h), 255)
    d = ImageDraw.Draw(img)
    y = 80
    for i, riga in enumerate(doc["righe"]):
        # « | » separa le colonne: la seconda comincia a metà pagina
        for j, pezzo in enumerate(riga.split(" | ")):
            d.text((90 + 600 * j, y), pezzo, fill=20, font=_font(40 if i == 0 else 30))
        y += 56
    if storta:
        img = img.rotate(rnd.uniform(-4, 4), expand=True, fillcolor=235)
        px = img.load()
        for _ in range(img.size[0] * img.size[1] // 60):
            x, yy = rnd.randrange(img.size[0]), rnd.randrange(img.size[1])
            px[x, yy] = max(0, px[x, yy] - rnd.randrange(40, 120))
        # luce che cala verso un angolo
        grad = Image.linear_gradient("L").resize(img.size).point(lambda v: 255 - v // 5)
        img = Image.composite(img, Image.new("L", img.size, 0), grad)
        img = img.filter(ImageFilter.GaussianBlur(0.8))
    return img.convert("RGB")


def degrada(img, seme: int = 1):
    """Una foto peggiore: risoluzione bassa (40 per cento), sfocata, JPEG molto compresso."""
    from PIL import ImageFilter
    small = img.resize((int(img.width * 0.4), int(img.height * 0.4)))
    small = small.filter(ImageFilter.GaussianBlur(0.9))
    buf = io.BytesIO()
    small.save(buf, "JPEG", quality=30)
    buf.seek(0)
    from PIL import Image
    return Image.open(buf).convert("RGB")


def _pdf_testo(doc: dict) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for riga in doc["righe"]:
        pdf.cell(0, 8, riga.replace("€", "EUR").replace(" | ", "   "), new_x="LMARGIN",
                 new_y="NEXT")
    return bytes(pdf.output())


def _pdf_scansione(doc: dict) -> bytes:
    from fpdf import FPDF
    img = immagine(doc, storta=False)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=70)
    buf.seek(0)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(buf, x=5, y=5, w=200)
    return bytes(pdf.output())


def _docx(doc: dict) -> bytes:
    import docx
    d = docx.Document()
    for riga in doc["righe"]:
        d.add_paragraph(riga)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def genera(cartella, formati=("pdf", "scansione", "foto", "docx")) -> list[dict]:
    """Scrive i documenti finti in `cartella` (solo quelli nei `formati` dati) e restituisce
    [{percorso, testo, tipo, formato, atteso}]."""
    out = []
    base = Path(cartella)
    base.mkdir(parents=True, exist_ok=True)
    for i, doc in enumerate(DOCUMENTI):
        if doc["formato"] not in formati:
            continue
        p = base / doc["file"]
        if doc["formato"] == "pdf":
            p.write_bytes(_pdf_testo(doc))
        elif doc["formato"] == "scansione":
            p.write_bytes(_pdf_scansione(doc))
        elif doc["formato"] == "foto":
            immagine(doc, storta=True, seme=i).save(p, "JPEG", quality=75)
        elif doc["formato"] == "docx":
            p.write_bytes(_docx(doc))
        out.append({"percorso": str(p), "testo": testo(doc), "tipo": doc["tipo"],
                    "formato": doc["formato"], "atteso": dict(doc["atteso"])})
    return out


# ─────────────────────────── per le prove: modello, OCR e persone finti ───────────────────────────
# La scheda che darebbe un modello bravo per ogni documento (con qualche valore inventato che il
# controllo contro il testo deve scartare), le sottocartelle personali, l'OCR che restituisce il
# testo vero, i profili Mario (amministra), Laura e Giulia.

try:
    import pypdfium2  # noqa: F401
    CON_PDF = True
except ImportError:
    CON_PDF = False

GREZZE = {
    "bolletta luce settembre 2026": ("bolletta", {
        "oggetto": "bolletta luce agosto 2026", "data_documento": "2026-09-10",
        "numero": "2026/0912345", "emittente": {"nome": "LUMEN ENERGIA S.p.A.",
                                                 "partita_iva": "01234567897"},
        "intestatari": [{"nome": "Mario Bianchi", "codice_fiscale": "BNCMRA80A01F205X"}],
        "importo_totale": 84.5, "scadenze": [{"data": "2026-10-15", "cosa": "pagamento"}],
        "indirizzi": ["Via dei Tigli 12, 20100 Milano"], "categoria": "luce",
        "periodo_dal": "2026-08-01", "periodo_al": "2026-08-31", "consumo": 210, "unita": "kWh",
        "codice_fornitura": "IT001E12345678"}),
    "luce 2025-09": ("bolletta", {
        "data_documento": "2025-09-09", "emittente": {"nome": "Lumen Energia SpA"},
        "intestatari": [{"nome": "BIANCHI MARIO", "codice_fiscale": "BNCMRA80A01F205X"}],
        "importo_totale": 97.2, "scadenze": [{"data": "2025-10-15", "cosa": "pagamento"}],
        "categoria": "luce", "consumo": 240, "codice_cliente": "C-999"}),     # inventato
    "gas ottobre": ("bolletta", {
        "data_documento": "2026-10-01", "numero": "G-778812",
        "emittente": {"nome": "GASNOVA S.R.L."}, "intestatari": [{"nome": "Mario Bianchi"}],
        "importo_totale": 62.3, "scadenze": [{"data": "2026-10-20", "cosa": "pagamento"}],
        "categoria": "gas", "codice_fornitura": "01234567890123", "consumo": 45, "unita": "Smc",
        "indirizzi": ["Via dei Tigli 12 - Milano"]}),
    "scontrino lavatrice": ("ricevuta", {
        "oggetto": "lavatrice Candor LV800", "data_documento": "2025-03-12", "numero": "118/2025",
        "emittente": {"nome": "ELETTRODOMESTICI ROSSI s.n.c.", "partita_iva": "09876543217"},
        "intestatari": [{"nome": "Mario Bianchi"}], "importo_totale": 479,
        "categoria": "elettrodomestici", "beni": [{"nome": "lavatrice Candor LV800"}],
        "voci": [{"descrizione": "Lavatrice Candor LV800", "importo": 449},
                 {"descrizione": "Trasporto e installazione", "importo": 30}]}),
    "garanzia lavatrice": ("garanzia", {
        "emittente": {"nome": "Elettrodomestici Rossi snc"}, "intestatari": [{"nome": "Mario Bianchi"}],
        "prodotto": "lavatrice Candor LV800", "marca": "Candor", "modello": "LV800",
        "numero_serie": "LV800-55A21", "data_acquisto": "2025-03-12", "durata_mesi": 24}),
    "polizza auto": ("assicurazione", {
        "data_documento": "2025-10-20", "numero": "AU-2026-445566",
        "emittente": {"nome": "AURORA ASSICURAZIONI S.p.A."}, "intestatari": [{"nome": "Mario Bianchi"}],
        "beni": [{"nome": "Fiat Panda", "identificativo": "AB123CD"}], "ramo": "auto",
        "data_inizio": "2025-11-01", "data_fine": "2026-11-01", "premio": 520,
        "scadenze": [{"data": "2026-12-01", "cosa": "rinnovo"}]}),               # inventata
    "referto Giulia": ("referto", {
        "emittente": {"nome": "POLIAMBULATORIO SAN LUCA"}, "intestatari": [{"nome": "Giulia Bianchi"}],
        "esame": "emocromo completo", "data_esame": "2026-09-05",
        "medico": {"nome": "Anna Verdi"}}),
    "carta identita Giulia": ("identita", {
        "numero": "CA12345AB", "emittente": {"nome": "Comune di Milano"},
        "intestatari": [{"nome": "Giulia Bianchi"}], "tipo_documento": "carta_identita",
        "data_nascita": "2012-07-14", "data_rilascio": "2021-07-14", "data_fine": "2031-07-14"}),
    "contratto internet": ("contratto", {
        "data_documento": "2026-01-20", "numero": "RV-99812",
        "emittente": {"nome": "RETE VELOCE S.p.A."}, "intestatari": [{"nome": "Mario Bianchi"}],
        "categoria": "internet", "canone": 29.9, "data_inizio": "2026-02-01", "durata_mesi": 24,
        "rinnovo_tacito": True, "preavviso_disdetta_giorni": 30, "periodicita": "mensile"}),
    "acqua settembre": ("bolletta", {
        "data_documento": "2026-09-05", "numero": "AN-55120",
        "emittente": {"nome": "ACQUE DEL NAVIGLIO S.p.A."}, "intestatari": [{"nome": "Laura Bianchi"}],
        "importo_totale": 79.81, "scadenze": [{"data": "2026-09-30", "cosa": "pagamento"}],
        "categoria": "acqua", "codice_cliente": "7781203"}),
}
CARTELLE = {"referto Giulia": "Giulia", "carta identita Giulia": "Giulia",
            "contratto internet": "Mario"}
TESTI = {}


def prepara_cartella(root: Path, tmp: Path) -> dict:
    """I documenti finti in `root`, con le sottocartelle personali (Giulia, Mario); senza
    pypdfium2 i PDF diventano testo (quelli con testo) e foto (le scansioni). Restituisce
    {nome senza estensione: testo vero}, che usa anche l'OCR finto."""
    docs = genera(tmp)
    for d in docs:
        p = Path(d["percorso"])
        TESTI[p.stem] = d["testo"]
        dest = root / CARTELLE.get(p.stem, "") / p.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        if p.suffix == ".pdf" and not CON_PDF:
            if d["formato"] == "pdf":
                dest = dest.with_suffix(".txt")
                dest.write_text(d["testo"], encoding="utf-8")
            else:
                immagine(next(x for x in DOCUMENTI if x["file"] == p.name)).save(
                    dest.with_suffix(".jpg"), "JPEG")
            continue
        shutil.copy(p, dest)
    return TESTI


class OcrFinto:
    def __init__(self):
        self.chiamate = 0

    def leggi(self, img, fonte=""):
        self.chiamate += 1
        stem = Path(fonte.split(",")[0]).stem
        return TESTI[stem]


class ClienteFinto:
    """Risponde come un modello bravo: il tipo e la scheda del documento (per nome del file)."""

    def __init__(self):
        self.chiamate = 0

    def chat(self, body, su_pezzo=None, controlla=None):
        self.chiamate += 1
        sistema = body["messages"][0]["content"]
        utente = body["messages"][1]["content"]
        nome = utente.split("\n", 1)[0].removeprefix("Nome del file: ")
        tipo, grezza = GREZZE[Path(nome).stem]
        if "Classifichi" in sistema:
            out = {"tipo": tipo}
        else:
            # tutti i campi dello schema, a null quelli che non ci sono (come gli output strutturati)
            out = {k: grezza.get(k, [] if body["format"]["properties"][k].get("type") == "array"
                                 else None) for k in body["format"]["properties"]}
        return {"content": json.dumps(out), "eval": 50, "prompt": 400, "s": 0.01}


class Prof:
    def __init__(self, pid, name, admin=False):
        self.id, self.name, self.admin = pid, name, admin


class Speakers:
    def __init__(self):
        self.p = {"Mario": Prof("mario-id", "Mario", True), "Laura": Prof("laura-id", "Laura"),
                  "Giulia": Prof("giulia-id", "Giulia")}

    def get(self, n):
        return self.p.get(n)

    def known_speakers(self):
        return list(self.p)


class SpeakerCtx:
    def __init__(self, name, level, grigia=False):
        self.current_speaker, self.current_level, self.from_session = name, level, grigia
        self.identified_by = "conversazione" if grigia else "voce"



if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    for d in genera(sys.argv[1] if len(sys.argv) > 1 else "archivio-finto"):
        print(d["formato"], d["percorso"])
