import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

"""Documenti a voce con il modello vero (Ollama), in una cartella temporanea.

Richieste come dette a voce → tool e argomenti giusti, file valido con i contenuti attesi
(riletto con le stesse librerie), risposta e tempi: prima frase (con l'attesa del file)
e generazione del JSON. Dario amministra, Bianca è familiare e non proprietaria del PC
(apre lo stesso i documenti suoi), Marco non ha documenti, più un ospite rifiutato e
qualche distrattore. Una lettera passa in secondo piano (attesa di 1 s): «te la preparo»,
annuncio a file pronto messo nella storia come in main.py, poi «sì, aprila». Il PC è
quello finto (prove/pc_finto.py): nessun file si apre davvero. Alla fine si misura quanto aspetta una domanda fatta mentre Ollama genera un
documento in secondo piano (il problema dell'«arbitro»).

    python prove\\prova_documenti_ollama.py        # 2 giri
    python prove\\prova_documenti_ollama.py 1      # 1 giro
"""

import re
import tempfile
import threading
import time
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ora_giusta import ora_o_tool  # noqa: E402  (03/10: l'ora giusta senza tool vale)
from calliope.agenda import Agenda
from calliope.brain import Brain
from calliope.config import Config
from calliope.documenti import Documenti
from calliope.documenti.consegna import LocalDelivery
from calliope.documenti.render import plain_text
from calliope.liste import Liste
from calliope.tools.builtin import build_registry
from calliope.tools.spec import ToolContext
from prove.pc_finto import FakePC

GIRI = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 2
errori = 0


def verifica(nome, ok, dettaglio=""):
    global errori
    errori += not ok
    print(f"{'ok ' if ok else 'ERR'} {nome}" + (f"  {dettaglio}" if dettaglio else ""), flush=True)


class Prof:
    def __init__(self, pid, name, admin=False):
        self.id, self.name, self.admin = pid, name, admin


class Speakers:
    def __init__(self):
        self.p = {"Dario": Prof("dario-id", "Dario", True), "Bianca": Prof("bianca-id", "Bianca"),
                  "Marco": Prof("marco-id", "Marco")}

    def get(self, n):
        return self.p.get(n)

    def known_speakers(self):
        return list(self.p)


class SpeakerCtx:
    def __init__(self, name, level):
        self.current_speaker, self.current_level, self.from_session = name, level, False


cfg = Config()
cfg_sfondo = Config()
cfg_sfondo.documenti_attesa_s = 1.0      # per far passare una lettera in secondo piano
IDS = {"Dario": "dario-id", "Bianca": "bianca-id", "Marco": "marco-id"}
DONE = re.compile(r"\b(ho preparato|è pronta|è pronto|te la preparo|te lo preparo)\b", re.I)


def ultimo(svc, chi):
    return svc.archive.last(IDS[chi]) if chi else None


def testo(svc, chi) -> str:
    d = ultimo(svc, chi)
    if not d or not Path(d["rif"]).exists():
        return ""
    return plain_text(d["formato"], Path(d["rif"]).read_bytes()).lower()


def foglio(svc, chi):
    from openpyxl import load_workbook
    d = ultimo(svc, chi)
    ws = load_workbook(d["rif"]).worksheets[0]
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def riga(rows, nome):
    return next((r for r in rows if r and str(r[0]).strip().lower().startswith(nome)), None)


def aperto(pc, prima, svc, chi):
    d = ultimo(svc, chi)
    return d is not None and pc.azioni[prima["n"]:] == [("apri", d["rif"])]


def somma_ok(rows, n):
    tot = riga(rows, "totale")
    return tot is not None and any(str(v).upper() == f"=SUM(B2:B{n + 1})" for v in tot[1:])


# (sessione, chi, livello, frase, tool atteso o None, controllo(risposta, argomenti, svc, pc, prima))
CASI = [
    ("lettera", "Dario", "amministra", "Preparami una lettera di disdetta per la palestra.",
     "documento_crea",
     lambda r, a, s, pc, p: ultimo(s, "Dario")["formato"] == "word"
     and "disdett" in testo(s, "Dario") and "palestra" in testo(s, "Dario")),
    ("lettera", "Dario", "amministra", "Cambia la data in 15 ottobre.", "documento_modifica",
     lambda r, a, s, pc, p: "15 ottobre" in testo(s, "Dario")
     and ultimo(s, "Dario")["versione"] == p["versione"] + 1),
    ("lettera", "Dario", "amministra", "Aprila.", "pc_apri_file",
     lambda r, a, s, pc, p: aperto(pc, p, s, "Dario")),
    ("spese", "Bianca", "familiare",
     "Fammi una tabella Excel con le spese di settembre: affitto 800, luce 90, gas 60, con il totale.",
     "documento_crea",
     lambda r, a, s, pc, p: ultimo(s, "Bianca")["formato"] == "excel"
     and [riga(foglio(s, "Bianca"), k)[1] for k in ("affitto", "luce", "gas")] == [800, 90, 60]
     and somma_ok(foglio(s, "Bianca"), 3)),
    ("spese", "Bianca", "familiare", "Aggiungi una riga: internet 30.", "documento_modifica",
     lambda r, a, s, pc, p: (riga(foglio(s, "Bianca"), "internet") or [0, 0])[1] == 30
     and somma_ok(foglio(s, "Bianca"), 4)),
    ("spese", "Bianca", "familiare", "Cambia l'importo della luce in 95.", "documento_modifica",
     lambda r, a, s, pc, p: riga(foglio(s, "Bianca"), "luce")[1] == 95
     and riga(foglio(s, "Bianca"), "affitto")[1] == 800
     and (riga(foglio(s, "Bianca"), "internet") or [0, 0])[1] == 30),
    ("spese", "Bianca", "familiare", "Aprilo.", "pc_apri_file",
     lambda r, a, s, pc, p: aperto(pc, p, s, "Bianca")),
    ("compiti", "Dario", "amministra",
     "Fai un PDF con l'elenco dei compiti di Matteo per domani: matematica pagina 40, "
     "leggere un capitolo del libro, ripassare storia.", "documento_crea",
     lambda r, a, s, pc, p: ultimo(s, "Dario")["formato"] == "pdf"
     and all(w in testo(s, "Dario") for w in ("40", "capitolo", "storia"))),
    ("compiti", "Dario", "amministra", "Aggiungi anche gli esercizi di inglese.",
     "documento_modifica",
     lambda r, a, s, pc, p: "inglese" in testo(s, "Dario") and "storia" in testo(s, "Dario")),
    ("invitati", "Bianca", "familiare",
     "Scrivimi un documento Word con la lista degli invitati alla festa: Marco, Giulia, Luca e Sara.",
     "documento_crea",
     lambda r, a, s, pc, p: ultimo(s, "Bianca")["formato"] == "word"
     and all(n in testo(s, "Bianca") for n in ("marco", "giulia", "luca", "sara"))),
    # In secondo piano (attesa di 1 s): «te la preparo», annuncio a file pronto, poi «sì»
    ("sfondo", "Bianca", "familiare",
     "Preparami una lettera di reclamo al comune per una buca pericolosa in via Roma.",
     "documento_crea",
     lambda r, a, s, pc, p: re.search(r"te la preparo", r, re.I) and "è pronta" in r
     and "buca" in testo(s, "Bianca")),
    ("sfondo", "Bianca", "familiare", "Sì, aprila.", "pc_apri_file",
     lambda r, a, s, pc, p: aperto(pc, p, s, "Bianca")),
    ("marco", "Marco", "familiare", "Cambia la data della lettera in 20 ottobre.", None,
     lambda r, a, s, pc, p: ultimo(s, "Marco") is None and "20 ottobre" not in testo(s, "Dario")
     and not re.search(r"\b(fatto|ho aggiornato|ho cambiato)\b", r, re.I)),
    ("ospite", None, "ospite", "Fammi una lettera per l'amministratore del condominio.", None,
     lambda r, a, s, pc, p: s.archive.last("ospite") is None and not DONE.search(r)
     and len(r) < 400),
    ("distrattori", "Dario", "amministra", "Che ore sono?", "ora_attuale",
     lambda r, a, s, pc, p: True),
    ("distrattori", "Dario", "amministra", "Quanto fa 800 più 90 più 60?", "calcola",
     lambda r, a, s, pc, p: "950" in r),
    ("distrattori", "Dario", "amministra", "Aggiungi il latte alla lista della spesa.",
     "lista_aggiungi", lambda r, a, s, pc, p: True),
]


def aspetta(svc, brain, limite=60.0) -> list[str]:
    """Aspetta i lavori in secondo piano e, come main.py, mette gli annunci nella storia."""
    t0 = time.perf_counter()
    while svc.busy() and time.perf_counter() - t0 < limite:
        time.sleep(0.05)
    msgs = []
    while not svc.done.empty():
        msg = svc.done.get_nowait()["messaggio"]
        brain.record_announcement(msg)
        msgs.append(msg)
    return msgs


righe = []
generazioni = []
for giro in range(1, GIRI + 1):
    tmp = Path(tempfile.mkdtemp(prefix="calliope-documenti-ollama-"))
    ag = Agenda(str(tmp / "memoria.db"))
    liste = Liste(str(tmp / "memoria.db"))
    pc = FakePC()
    pcs = {"portatile": pc}
    svc = Documenti(cfg, str(tmp / "memoria.db"), delivery=LocalDelivery(tmp / "Documenti"),
                    pcs=pcs)
    # Con i tool degli schermi, degli agenti e dell'ufficio, come in Calliope (schermi accesi
    # di predefinito e agenti con dgx.yaml, 02/10; ufficio acceso di predefinito, 03/10)
    from calliope.ufficio import Ufficio
    uff = Ufficio(cfg, str(tmp / "memoria.db"), svc)
    reg = build_registry(pc=pcs, documenti=svc.formati, schermi=True, agenti=True,
                         ufficio=uff.nomi_modelli())
    brains = {}
    if giro == 1 and svc.writer._native is not None:
        # Solo con Ollama: con il backend openai (CALLIOPE_LLM_PROFILO, 03/10) il server
        # tiene già il modello in memoria
        svc.writer._native.warmup()
    for sessione, chi, livello, frase, atteso, controllo in CASI:
        key = (sessione, chi)
        if key not in brains:
            ctx = ToolContext(cfg=cfg_sfondo if sessione == "sfondo" else cfg, speakers=Speakers(),
                              speaker_ctx=SpeakerCtx(chi, livello),
                              speaker=None, agenda=ag, liste=liste, pc=pcs, documenti=svc)
            brains[key] = Brain(cfg, reg, ctx)
        b = brains[key]
        last = ultimo(svc, chi)
        prima = {"n": len(pc.azioni), "versione": last["versione"] if last else 0}
        n_stats = len(svc.stats)
        t0 = time.perf_counter()
        primo, parti = None, []
        for pezzo in b.stream_reply(frase, livello):
            if primo is None and pezzo.strip():
                primo = time.perf_counter() - t0
            parti.append(pezzo)
        totale = time.perf_counter() - t0
        annunci = aspetta(svc, b)
        attesa_bg = time.perf_counter() - t0
        risposta = " ".join(["".join(parti).strip()] + annunci).strip()
        tools = [t["nome"] for t in b.last_tools]
        args = next((t["argomenti"] for t in reversed(b.last_tools) if t["nome"] == atteso), {})
        if atteso is None:
            ok_tool = not any(t["nome"].startswith("documento_") and t["ok"] for t in b.last_tools)
        else:
            ok_tool = atteso in tools or ora_o_tool(atteso, tools, risposta)
        try:
            ok_stato = bool(controllo(risposta, args, svc, pc, prima))
        except Exception as e:                      # un controllo rotto è un errore, non un crash
            ok_stato, risposta = False, f"{risposta} [controllo: {type(e).__name__}: {e}]"
        gen = svc.stats[n_stats:]
        for g in gen:
            generazioni.append({**g, "giro": giro, "frase": frase,
                                "secondo_piano": bool(annunci)})
        argomenti = "; ".join(f"{t['nome']}({', '.join(f'{k}={v!r}' for k, v in t['argomenti'].items())})"
                              for t in b.last_tools) or "—"
        tempi = f"{primo or 0:.2f}/{totale:.2f}s" + (f" (file a {attesa_bg:.1f}s)" if annunci else "")
        if gen:
            tempi += f" gen {gen[-1]['generazione_s']:.2f}s {gen[-1].get('token')} tok"
        verifica(f"[{giro}] {chi or 'ospite'}: «{frase}» → {argomenti}", ok_tool and ok_stato,
                 f"{tempi}  {risposta[:150]!r}")
        righe.append((giro, frase, ok_tool and ok_stato, primo or 0, totale))
    svc.close()

# ── L'arbitro: una domanda mentre Ollama genera un documento in secondo piano ──
tmp = Path(tempfile.mkdtemp(prefix="calliope-documenti-arbitro-"))
svc = Documenti(cfg, str(tmp / "memoria.db"), delivery=LocalDelivery(tmp / "Documenti"))
reg = build_registry(documenti=svc.formati, schermi=True, agenti=True)
ctx = ToolContext(cfg=cfg, speakers=Speakers(), speaker_ctx=SpeakerCtx("Dario", "amministra"),
                  speaker=None, documenti=svc)


def prima_frase(frase):
    b = Brain(cfg, reg, ctx)
    t0 = time.perf_counter()
    for pezzo in b.stream_reply(frase, "amministra"):
        if pezzo.strip():
            return time.perf_counter() - t0
    return time.perf_counter() - t0


libera = [prima_frase("Che ore sono?") for _ in range(2)]
occupata = []
for _ in range(2):
    job = svc.crea("dario-id", "Dario", "word", "lettera di reclamo al comune per una buca in "
                   "via Roma", "")
    time.sleep(0.5)                         # la generazione è partita
    occupata.append(prima_frase("Che ore sono?"))
    job.event.wait(60)
svc.close()
print(f"\narbitro: «che ore sono?» con Ollama libero {min(libera):.2f}s, mentre genera una "
      f"lettera {', '.join(f'{x:.2f}' for x in occupata)}s")

prime = sorted(r[3] for r in righe)
print(f"prima frase: mediana {prime[len(prime) // 2]:.2f}s, massimo {prime[-1]:.2f}s; "
      f"riuscite {sum(r[2] for r in righe)}/{len(righe)}")
for kind in ("word", "excel", "pdf"):
    for job in ("crea", "modifica"):
        g = sorted(x["generazione_s"] for x in generazioni if x["formato"] == kind and x["lavoro"] == job)
        if g:
            bg = sum(x["secondo_piano"] for x in generazioni if x["formato"] == kind and x["lavoro"] == job)
            print(f"generazione {kind} {job}: mediana {g[len(g) // 2]:.2f}s, "
                  f"da {g[0]:.2f} a {g[-1]:.2f}s ({len(g)} volte, {bg} in secondo piano)")
print(f"{errori} errori" if errori else "Tutto a posto.")
sys.exit(1 if errori else 0)
