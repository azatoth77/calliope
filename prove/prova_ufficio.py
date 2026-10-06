import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

"""Prova a secco dell'ufficio (calliope/ufficio/, calliope/tools/ufficio.py).

Niente Ollama: l'estrazione dei campi e lo scrittore sono finti (risposte preparate). Tutto
in una cartella temporanea. Si provano: partita IVA e codice fiscale; i conti (più aliquote,
arrotondamenti, IVA inclusa, ritenuta, cassa, bollo, forfettario, esenzioni ed esportazioni);
l'XML FatturaPA (controlli interni e, se lo schema è in fatturapa/, l'XSD ufficiale, anche
contro un XML rovinato); la numerazione (sequenza, thread e processi in concorrenza, nessun
buco, un errore che non consuma il numero, ordine delle date, annullo con la nota); la
rubrica (ricerca, ambito personale, doppioni); i modelli dell'utente in Word e PowerPoint
(creati qui, con segnaposti, righe e diapositive ripetute, valori con «&» e «<»); il flusso
a voce (domande sui dati mancanti, proposta, «sì» nel turno dopo, permessi, PDF e XML
scritti); i tool dal registro, il prompt e quanto allungano il prefisso.
"""

import datetime
import io
import json
import multiprocessing
import tempfile
import threading
import time
import xml.etree.ElementTree as ET
from decimal import Decimal
from pathlib import Path

from calliope.config import Config
from calliope.documenti import Documenti
from calliope.documenti.consegna import LocalDelivery
from calliope.documenti.render import plain_text
from calliope.tools.builtin import build_registry
from calliope.tools.spec import ToolContext
from calliope.ufficio import conti, fatturapa
from calliope.ufficio.modelli import carica, compila_docx, compila_pptx, librerie, risolvi
from calliope.ufficio.numerazione import NumerazioneError, Numeratore
from calliope.ufficio.rubrica import Rubrica, codice_fiscale_ok, controlla, partita_iva_ok
from calliope.ufficio.servizio import Ufficio, parse_data

errori = 0


def verifica(nome, ok, dettaglio=""):
    global errori
    errori += not ok
    print(f"{'ok ' if ok else 'ERR'} {nome}" + (f"  {dettaglio}" if dettaglio else ""))


TMP = Path(tempfile.mkdtemp(prefix="calliope-ufficio-"))
XSD = fatturapa.cartella_xsd(Config())
OGGI = datetime.date(2026, 10, 2)
D = Decimal

EMITTENTE = {"denominazione": "Esempio Srl", "partita_iva": "01234567897",
             "codice_fiscale": "01234567897", "regime_fiscale": "RF01", "indirizzo": "Via Roma",
             "civico": "3", "cap": "20121", "comune": "Milano", "provincia": "MI",
             "iban": "IT60 X054 2811 1010 0000 0123 456", "ritenuta": "20",
             "giorni_pagamento": "30"}
ROSSI = {"tipo": "cliente", "denominazione": "Rossi & Figli S.r.l.", "partita_iva": "12345678903",
         "indirizzo": "Via Verdi", "civico": "5", "cap": "00184", "comune": "Roma",
         "provincia": "RM", "pec": "rossi@pec.it"}
PRIVATO = {"tipo": "cliente", "nome": "Mario", "cognome": "Rossi",
           "codice_fiscale": "RSSMRA80A01H501U", "indirizzo": "Via Po", "civico": "1",
           "cap": "10121", "comune": "Torino", "provincia": "TO"}


# ─────────────────────────── codici ───────────────────────────

def prova_codici():
    verifica("partita IVA valida", partita_iva_ok("01234567897") and partita_iva_ok("12345678903"))
    verifica("partita IVA con la cifra di controllo sbagliata", not partita_iva_ok("01234567890"))
    verifica("partita IVA corta o con lettere", not partita_iva_ok("1234567") and
             not partita_iva_ok("0123456789A"))
    verifica("codice fiscale di persona valido", codice_fiscale_ok("RSSMRA80A01H501U"))
    verifica("codice fiscale con il controllo sbagliato", not codice_fiscale_ok("RSSMRA80A01H501X"))
    verifica("codice fiscale numerico (società) = partita IVA", codice_fiscale_ok("01234567897"))
    d, e = controlla({"partita_iva": "IT 012 345 678 97", "provincia": "milano", "cap": "2012",
                      "email": "mario chiocciola rossi punto it", "codice_destinatario": "abc"})
    verifica("forma normalizzata: spazi, IT, provincia per nome, email detta",
             d["partita_iva"] == "01234567897" and d["provincia"] == "MI"
             and d["email"] == "mario@rossi.it", str(d))
    verifica("errori di forma: CAP corto e codice destinatario", set(e) == {"cap",
                                                                          "codice_destinatario"},
             str(e))
    from calliope.ufficio.servizio import iban_ok
    verifica("IBAN con il controllo", iban_ok("IT60X0542811101000000123456")
             and not iban_ok("IT61X0542811101000000123456"))
    verifica("date dette", parse_data("oggi", OGGI) == OGGI and parse_data("15 ottobre", OGGI)
             == datetime.date(2026, 10, 15) and parse_data("1/9/26", OGGI)
             == datetime.date(2026, 9, 1) and parse_data("ieri", OGGI) == datetime.date(2026, 10, 1)
             and parse_data("boh", OGGI) is None)


# ─────────────────────────── conti ───────────────────────────

def righe(*r, **kw):
    return conti.prepara_righe([dict(x) for x in r], kw.pop("aliq", 22), **kw)


def prova_conti():
    t = conti.calcola(righe({"descrizione": "Consulenza", "quantita": 2, "prezzo": 500},
                            ritenuta_predefinita=True), ritenuta={"aliquota": 20})
    verifica("2 × 500 al 22% con ritenuta 20%",
             (t["imponibile"], t["imposta"], t["totale"], t["ritenuta"]["importo"],
              t["da_pagare"]) == (D("1000.00"), D("220.00"), D("1220.00"), D("200.00"),
                                  D("1020.00")), conti.riassunto_detto(t))
    verifica("riassunto detto", conti.riassunto_detto(t) == "imponibile 1.000 euro, IVA 220 "
             "euro, totale 1.220 euro, ritenuta 200 euro, da pagare 1.020 euro")
    t = conti.calcola(righe({"descrizione": "A", "prezzo": 100}, {"descrizione": "B",
                            "prezzo": 50, "aliquota": 10}, {"descrizione": "C", "prezzo": 30,
                                                             "aliquota": 4}))
    verifica("più aliquote: un riepilogo per aliquota, dalla più alta",
             [(e["aliquota"], e["imponibile"], e["imposta"]) for e in t["riepilogo"]]
             == [(D(22), D("100.00"), D("22.00")), (D(10), D("50.00"), D("5.00")),
                 (D(4), D("30.00"), D("1.20"))] and t["totale"] == D("208.20"))
    # 10,75 × 22% = 2,365: arrotondamento commerciale 2,37 (quello «del banchiere» darebbe 2,36)
    t = conti.calcola(righe({"descrizione": "A", "prezzo": "10,75"}))
    verifica("arrotondamento commerciale (2,365 → 2,37)", t["imposta"] == D("2.37"))
    # IVA arrotondata per aliquota, non per riga: 3 × 0,15 al 22% = 0,099 → 0,10 (per riga 0,09)
    t = conti.calcola(righe(*[{"descrizione": f"R{i}", "prezzo": "0.15"} for i in range(3)]))
    verifica("IVA arrotondata una volta per aliquota", t["imposta"] == D("0.10"), str(t["imposta"]))
    t = conti.calcola(righe({"descrizione": "A", "quantita": 3, "prezzo": "10.33"}))
    verifica("quantità × prezzo: 3 × 10,33 = 30,99", t["righe"][0]["totale"] == D("30.99"))
    for tot, aliq in (("100", 22), ("99.99", 22), ("1234.56", 10), ("7", 4)):
        t = conti.calcola(righe({"descrizione": "A", "prezzo": tot, "aliquota": aliq},
                                iva_inclusa=True))
        verifica(f"IVA inclusa {tot} al {aliq}%: il totale resta quello detto",
                 t["totale"] == D(tot).quantize(D("0.01")), f"{t['totale']} (unitario "
                                                           f"{t['righe'][0]['prezzo']})")
    t = conti.calcola(righe({"descrizione": "Sito", "prezzo": 500}, aliq=0, natura_predefinita=
                            "N2.2"), bollo="auto", bollo_addebito=True)
    verifica("forfettario 500: N2.2 con il riferimento, bollo addebitato come riga N1",
             [(e["natura"], e["imponibile"]) for e in t["riepilogo"]] == [("N2.2", D("500.00")),
                                                                         ("N1", D("2.00"))]
             and t["bollo"] == D("2.00") and t["totale"] == D("502.00")
             and "190/2014" in t["riepilogo"][0]["riferimento"])
    t = conti.calcola(righe({"descrizione": "Sito", "prezzo": "77.47"}, aliq=0,
                            natura_predefinita="N2.2"), bollo_addebito=True)
    verifica("bollo solo sopra 77,47", t["bollo"] is None and t["totale"] == D("77.47"))
    t = conti.calcola(righe({"descrizione": "A", "prezzo": 100}, {"descrizione": "Corso",
                            "prezzo": 80, "natura": "N4"}), bollo_addebito=False)
    verifica("esente N4 sopra soglia: bollo dovuto, a carico dell'emittente",
             t["bollo"] == D("2.00") and not t["bollo_addebitato"] and t["totale"] == D("202.00"))
    t = conti.calcola(righe({"descrizione": "Export", "prezzo": 1000, "natura": "N3.1"}))
    verifica("esportazione N3.1: niente bollo", t["bollo"] is None)
    t = conti.calcola(righe({"descrizione": "Consulenza", "prezzo": 1000},
                            ritenuta_predefinita=True), ritenuta={"aliquota": 20},
                      cassa={"tipo": "TC22", "aliquota": 4, "aliquota_iva": 22})
    verifica("cassa 4% (rivalsa INPS) con IVA e ritenuta solo sul compenso",
             (t["cassa"]["importo"], t["imponibile"], t["imposta"], t["totale"],
              t["ritenuta"]["importo"], t["da_pagare"]) == (
                 D("40.00"), D("1040.00"), D("228.80"), D("1268.80"), D("200.00"),
                 D("1068.80")), conti.riassunto_detto(t))
    for bad, why in (([{"descrizione": "A", "prezzo": 10, "aliquota": 21}], "aliquota 21"),
                     ([{"descrizione": "A", "prezzo": 10, "aliquota": 0}], "IVA 0 senza natura"),
                     ([{"descrizione": "A", "prezzo": 10, "natura": "N2"}], "natura generica N2"),
                     ([{"descrizione": "", "prezzo": 10}], "descrizione vuota"),
                     ([{"descrizione": "A", "prezzo": None}], "prezzo mancante"),
                     ([{"descrizione": "A", "prezzo": 10, "quantita": 0}], "quantità zero"),
                     ([], "nessuna riga")):
        try:
            conti.prepara_righe(bad, 22)
            verifica(f"rifiutato: {why}", False)
        except conti.DatiNonValidi as e:
            verifica(f"rifiutato: {why}", True, e.errori[0])


# ─────────────────────────── XML FatturaPA ───────────────────────────

def fattura_xml(cliente, rr, numero="1/2026", tipo="TD01", **kw):
    t = conti.calcola(rr, kw.pop("ritenuta", None), kw.pop("cassa", None),
                      bollo_addebito=kw.pop("bollo_addebito", True))
    em = controlla(EMITTENTE)[0] | {"regime_fiscale": kw.pop("regime", "RF01")}
    f = {"tipo": tipo, "numero": numero, "data": OGGI, "emittente": em,
         "trasmittente": "01234567897", "cliente": controlla(cliente)[0],
         "progressivo": fatturapa.progressivo("fatture", 1, 2026),
         "causale": kw.pop("causale", ""), "pagamento": kw.pop("pagamento", None), **kw}
    return f, t, fatturapa.costruisci(f, t)


def campo(xml: bytes, percorso: str) -> list[str]:
    root = ET.fromstring(xml)
    return [e.text for e in root.iter() if e.tag.split("}")[-1] == percorso]


def prova_xml():
    casi = {
        "ordinaria con ritenuta e pagamento": fattura_xml(
            ROSSI, righe({"descrizione": "Consulenza “sito” – € 500", "quantita": 2,
                          "prezzo": 500}, ritenuta_predefinita=True),
            ritenuta={"aliquota": 20, "tipo": "RT02"},
            pagamento={"modalita": "MP05", "iban": "IT60X0542811101000000123456",
                       "scadenza": OGGI + datetime.timedelta(days=30)}),
        "più aliquote ed esente N4 con bollo": fattura_xml(
            ROSSI, righe({"descrizione": "A", "prezzo": 100}, {"descrizione": "Libro",
                         "prezzo": 30, "aliquota": 4}, {"descrizione": "Corso", "prezzo": 80,
                                                         "natura": "N4"})),
        "forfettario a un privato (solo codice fiscale)": fattura_xml(
            PRIVATO, righe({"descrizione": "Sito vetrina", "prezzo": 900}, aliq=0,
                           natura_predefinita="N2.2"), regime="RF19"),
        "non soggetta N2.1 e cassa": fattura_xml(
            ROSSI, righe({"descrizione": "Servizio a cliente UE", "prezzo": 300,
                          "natura": "N2.1"}, {"descrizione": "Consulenza", "prezzo": 1000},
                         ritenuta_predefinita=True),
            ritenuta={"aliquota": 20}, cassa={"tipo": "TC22", "aliquota": 4,
                                              "aliquota_iva": 22}),
        "nota di credito TD04": fattura_xml(
            ROSSI, righe({"descrizione": "Storno parziale", "prezzo": 100}), numero="NC1/2026",
            tipo="TD04", collegata={"numero": "1/2026", "data": OGGI}),
        "causale lunga (più di 200 caratteri)": fattura_xml(
            ROSSI, righe({"descrizione": "A", "prezzo": 10}), causale="x" * 450),
    }
    xsd = fatturapa.xsd_presente(XSD)
    if not xsd:
        print(f"   (schema XSD non presente in {XSD}: solo i controlli interni; "
              f"python -m calliope.ufficio --scarica-xsd)")
    for nome, (f, t, xml) in casi.items():
        err = fatturapa.controlla(f, t)
        verifica(f"XML {nome}: controlli interni", not err, "; ".join(err))
        if xsd:
            e = fatturapa.valida_xsd(xml, XSD)
            verifica(f"XML {nome}: valido per l'XSD ufficiale", e == [], "; ".join(e or []))
    f, t, xml = casi["ordinaria con ritenuta e pagamento"]
    verifica("XML: totale, ritenuta, pagamento al netto della ritenuta",
             campo(xml, "ImportoTotaleDocumento") == ["1220.00"]
             and campo(xml, "ImportoRitenuta") == ["200.00"]
             and campo(xml, "ImportoPagamento") == ["1020.00"]
             and campo(xml, "TipoRitenuta") == ["RT02"]
             and campo(xml, "IBAN") == ["IT60X0542811101000000123456"])
    verifica("XML: testo in Latin-1 (€ → EUR, virgolette semplici)",
             campo(xml, "Descrizione") == ['Consulenza "sito" - EUR 500'], campo(xml, "Descrizione")[0])
    verifica("XML: & nel nome del cliente scritto bene", b"Rossi &amp; Figli" in xml)
    verifica("XML: PEC del cliente con codice destinatario 0000000",
             campo(xml, "CodiceDestinatario") == ["0000000"] and campo(xml, "PECDestinatario")
             == ["rossi@pec.it"])
    _, _, xml = casi["più aliquote ed esente N4 con bollo"]
    verifica("XML: tre riepiloghi, N4 con il riferimento normativo, bollo virtuale",
             campo(xml, "AliquotaIVA")[-4:] == ["22.00", "4.00", "0.00", "0.00"]
             and "N4" in campo(xml, "Natura") and campo(xml, "BolloVirtuale") == ["SI"]
             and any("art. 10" in r for r in campo(xml, "RiferimentoNormativo")))
    _, _, xml = casi["forfettario a un privato (solo codice fiscale)"]
    verifica("XML: privato senza IdFiscaleIVA, con il codice fiscale, RF19",
             campo(xml, "CodiceFiscale")[-1] == "RSSMRA80A01H501U"
             and campo(xml, "RegimeFiscale") == ["RF19"] and len(campo(xml, "IdFiscaleIVA")) == 1)
    _, _, xml = casi["causale lunga (più di 200 caratteri)"]
    verifica("XML: causale lunga divisa in pezzi da 200", [len(c) for c in campo(xml, "Causale")]
             == [200, 200, 50])
    _, _, xml = casi["nota di credito TD04"]
    verifica("XML: nota di credito con la fattura collegata",
             campo(xml, "TipoDocumento") == ["TD04"] and campo(xml, "IdDocumento") == ["1/2026"])
    verifica("nome del file e progressivo", fatturapa.nome_file("01234567897", fatturapa.progressivo(
        "fatture", 1, 2026)) == "IT01234567897_" + fatturapa.progressivo("fatture", 1, 2026) + ".xml"
             and len({fatturapa.progressivo(s, n, a) for s in ("fatture", "note_credito")
                      for n in (1, 2, 49999) for a in (2026, 2027)}) == 12
             and all(len(fatturapa.progressivo("fatture", n, 2099)) == 5 for n in (1, 49999)))
    f, t, _ = fattura_xml(ROSSI, righe({"descrizione": "A", "prezzo": 10}))
    f["cliente"] = dict(f["cliente"], partita_iva="12345678900")
    f["numero"] = "senza cifre"
    err = fatturapa.controlla(f, t)
    verifica("controlli interni: partita IVA del cliente e numero senza cifre",
             any("cliente non è valida" in e for e in err) and any("cifra" in e for e in err),
             "; ".join(err))
    if xsd:
        _, _, xml = fattura_xml(ROSSI, righe({"descrizione": "A", "prezzo": 10}))
        rotto = xml.replace(b"<Divisa>EUR</Divisa>", b"").replace(b"<CAP>00184</CAP>",
                                                                   b"<CAP>184</CAP>")
        e = fatturapa.valida_xsd(rotto, XSD)
        verifica("XSD: un XML rovinato (Divisa mancante, CAP corto) non passa", bool(e),
                 (e or [""])[0][:90])


# ─────────────────────────── numerazione ───────────────────────────

def _processo(db, n, out):
    num = Numeratore(db)
    got = []
    for _ in range(n):
        got.append(num.emetti("fatture", OGGI, lambda k, s: {"file": []})[0])
    out.put(got)


def prova_numerazione():
    db = str(TMP / "numeri.db")
    num = Numeratore(db)
    nums = [num.emetti("fatture", OGGI, lambda k, s: {"file": []})[0] for _ in range(3)]
    verifica("sequenza 1, 2, 3", nums == [1, 2, 3])
    verifica("formato del numero", num.testo("fatture", 3, 2026) == "3/2026"
             and num.testo("ddt", 4, 2026) == "DDT4/2026")
    try:
        num.emetti("fatture", OGGI, lambda k, s: (_ for _ in ()).throw(RuntimeError("guasto")))
    except RuntimeError:
        pass
    verifica("un errore nel preparare il file non consuma il numero",
             num.prossimo("fatture", 2026) == 4)
    try:
        num.emetti("fatture", OGGI - datetime.timedelta(days=5), lambda k, s: {"file": []})
        verifica("data prima dell'ultima fattura rifiutata", False)
    except NumerazioneError as e:
        verifica("data prima dell'ultima fattura rifiutata", True, str(e)[:70])
    verifica("anno nuovo: si riparte da 1", num.emetti(
        "fatture", datetime.date(2027, 1, 2), lambda k, s: {"file": []})[0] == 1)
    # Thread in concorrenza
    got, lock = [], threading.Lock()

    def t():
        for _ in range(10):
            n = Numeratore(db).emetti("preventivi", OGGI, lambda k, s: (time.sleep(0.001),
                                                                         {"file": []})[1])[0]
            with lock:
                got.append(n)
    ths = [threading.Thread(target=t) for _ in range(8)]
    for th in ths:
        th.start()
    for th in ths:
        th.join()
    verifica("8 thread × 10: 80 numeri diversi da 1 a 80, nessun buco",
             sorted(got) == list(range(1, 81)) and num.buchi("preventivi", 2026) == [])
    # Processi diversi sullo stesso file (BEGIN IMMEDIATE)
    ctx = multiprocessing.get_context("spawn")
    q = ctx.Queue()
    procs = [ctx.Process(target=_processo, args=(db, 15, q)) for _ in range(3)]
    for p in procs:
        p.start()
    tutti = [n for _ in procs for n in q.get(timeout=60)]
    for p in procs:
        p.join(30)
    attesi = list(range(4, 4 + 45))
    verifica("3 processi × 15 sullo stesso database: nessun doppione né buco",
             sorted(tutti) == attesi and num.buchi("fatture", 2026) == [],
             f"{len(tutti)} numeri")
    try:
        num.annulla("fatture", 2026, 2, "")
        verifica("annullo senza nota rifiutato", False)
    except NumerazioneError:
        verifica("annullo senza nota rifiutato", True)
    num.annulla("fatture", 2026, 2, "preparata per errore, mai inviata")
    e = [d for d in num.elenco("fatture", 2026) if d["numero"] == 2][0]
    verifica("annullo con la nota: il numero resta, segnato annullato",
             e["stato"] == "annullato" and "mai inviata" in e["nota"]
             and num.prossimo("fatture", 2026) == 49 and num.buchi("fatture", 2026) == [])
    try:
        num.annulla("fatture", 2026, 2, "di nuovo per errore")
        verifica("annullare due volte rifiutato", False)
    except NumerazioneError:
        verifica("annullare due volte rifiutato", True)


# ─────────────────────────── rubrica ───────────────────────────

def prova_rubrica():
    r = Rubrica(str(TMP / "rubrica.db"))
    r.aggiungi(ROSSI, "casa", "dario-id")
    r.aggiungi({"denominazione": "Rossetti Spa", "partita_iva": "01234567897"}, "casa", "x")
    r.aggiungi(PRIVATO, "bianca-id", "bianca-id")
    c, alt = r.trova("Rossi e figli", "dario-id")
    verifica("«Rossi e figli» trova Rossi & Figli S.r.l.", c is not None and c["partita_iva"]
             == "12345678903", str(alt))
    c, alt = r.trova("rossi srl", "dario-id")
    verifica("«rossi srl» trova la ditta (senza la forma societaria)", c is not None)
    c, _ = r.trova("Mario Rossi", "dario-id")
    verifica("il contatto personale di Bianca non si vede da Dario",
             c is None or c.get("ambito") == "casa")
    c, _ = r.trova("Mario Rossi", "bianca-id")
    verifica("Bianca vede il suo", c is not None and c["codice_fiscale"] == "RSSMRA80A01H501U")
    c, _ = r.trova("12345678903", "dario-id")
    verifica("si trova anche per partita IVA", c is not None)
    verifica("doppione per partita IVA", r.duplicato({"partita_iva": "12345678903"}, "dario-id")
             is not None)
    try:
        r.aggiungi({"denominazione": "X", "partita_iva": "123"}, "casa", "x")
        verifica("partita IVA sbagliata non entra in rubrica", False)
    except ValueError:
        verifica("partita IVA sbagliata non entra in rubrica", True)


# ─────────────────────────── modelli Word e PowerPoint ───────────────────────────

def crea_modelli(cartella: Path):
    from docx import Document
    d = Document()
    d.add_heading("Preventivo {{ numero }}", 1)
    d.add_paragraph("Spett.le {{ cliente.nome }}, {{ cliente.indirizzo_completo }}")
    d.add_paragraph("Oggetto: {{ oggetto }}")
    d.add_paragraph("{{ descrizione }}")
    tb = d.add_table(rows=3, cols=2)
    tb.cell(0, 0).text, tb.cell(0, 1).text = "{%tr for v in voci %}", ""
    tb.cell(1, 0).text, tb.cell(1, 1).text = "{{ v.descrizione }}", "{{ v.importo }}"
    tb.cell(2, 0).text, tb.cell(2, 1).text = "{%tr endfor %}", ""
    d.add_paragraph("Totale {{ totali.totale }}")
    d.save(cartella / "Offerta sito.docx")
    (cartella / "Offerta sito.yaml").write_text(
        "nome: offerta sito\nalias: [preventivo sito]\ntitolo: Offerta per un sito\n"
        "serie: preventivi\ncampi:\n"
        "  cliente: {tipo: anagrafica, descrizione: il cliente}\n"
        "  oggetto: {tipo: testo, descrizione: oggetto}\n"
        "  descrizione: {tipo: testo_lungo, obbligatorio: false, descrizione: il lavoro}\n"
        "  voci: {tipo: righe, descrizione: le voci}\n", encoding="utf-8")
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Riunione {{ titolo }}"
    s.placeholders[1].text = "Del {{ data }}"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "{# ripeti punti come p #}Punto: {{ p }}"
    s.placeholders[1].text = "Riunione {{ titolo }}"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Partecipanti"
    tbl = s.shapes.add_table(2, 1, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tbl.cell(0, 0).text = "Nome"
    tbl.cell(1, 0).text = "{# riga partecipanti come x #}{{ x }}"
    prs.save(cartella / "Verbale.pptx")
    (cartella / "Verbale.yaml").write_text(
        "nome: verbale\ntitolo: Verbale di riunione\ncampi:\n"
        "  titolo: {tipo: testo}\n  punti: {tipo: elenco}\n  partecipanti: {tipo: elenco}\n",
        encoding="utf-8")
    (cartella / "Rotto.yaml").write_text("nome: rotto\ncampi: {}\n", encoding="utf-8")
    (cartella / "fattura.yaml").write_text("nome: fattura\ncampi: {x: {tipo: testo}}\n",
                                           encoding="utf-8")
    from docx import Document as D2
    D2().save(cartella / "fattura.docx")


def prova_modelli():
    lib = librerie()
    if not (lib["docx"] and lib["pptx"]):
        print(f"SALTATA IN PARTE: docxtpl o python-pptx mancanti: {lib}; modelli Word e PowerPoint saltati")
        return
    cartella = TMP / "Modelli"
    cartella.mkdir()
    crea_modelli(cartella)
    cat, err = carica(cartella)
    verifica("catalogo: 4 pronti + 2 dell'utente", sorted(cat) == sorted(
        ["fattura", "nota di credito", "preventivo", "ddt", "offerta sito", "verbale"]),
             str(sorted(cat)))
    verifica("modelli rovinati e nome riservato segnalati",
             any(e.startswith("rotto") for e in err) and any("riservato" in e for e in err),
             str(err))
    verifica("nome detto con alias", risolvi(cat, "il preventivo sito").nome == "offerta sito"
             and risolvi(cat, "la bolla").nome == "ddt")
    m = cat["offerta sito"]
    verifica("campi mancanti dal modello", m.mancanti({"cliente": "x"}) == ["oggetto", "voci"])
    ctx = {"numero": "P1/2026", "cliente": {"nome": "Bar <Sole> & Luna",
                                            "indirizzo_completo": "Via Roma 1, Milano"},
           "oggetto": "Sito", "descrizione": "Cinque pagine.",
           "voci": [{"descrizione": "Grafica", "importo": "500,00 €"},
                    {"descrizione": "Sviluppo", "importo": "1.000,00 €"}],
           "totali": {"totale": "1.830,00 €"}}
    data = compila_docx(m.file, ctx)
    txt = plain_text("word", data)
    verifica("docx: segnaposti, righe ripetute, «&» e «<» nei valori",
             "Bar <Sole> & Luna" in txt and "Grafica | 500,00 €" in txt
             and "Sviluppo | 1.000,00 €" in txt and "{{" not in txt and "{%" not in txt
             and "Totale 1.830,00 €" in txt, txt[:200].replace("\n", " / "))
    from pptx import Presentation
    data = compila_pptx(cat["verbale"].file, {"titolo": "di ottobre", "data": "2 ottobre 2026",
                                              "punti": ["Budget", "Sito", "Varie"],
                                              "partecipanti": ["Dario", "Bianca"]})
    prs = Presentation(io.BytesIO(data))
    titoli = [s.shapes.title.text for s in prs.slides]
    verifica("pptx: diapositiva ripetuta per ogni punto, nell'ordine",
             titoli == ["Riunione di ottobre", "Punto: Budget", "Punto: Sito", "Punto: Varie",
                        "Partecipanti"], str(titoli))
    tabella = [sh for sh in prs.slides[4].shapes if sh.has_table][0].table
    verifica("pptx: riga della tabella ripetuta", [r.cells[0].text for r in tabella.rows]
             == ["Nome", "Dario", "Bianca"])
    resto = " ".join(p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
                     for p in sh.text_frame.paragraphs)
    verifica("pptx: nessun segnaposto rimasto", "{{" not in resto and "{#" not in resto)
    data = compila_pptx(cat["verbale"].file, {"titolo": "x", "data": "y", "punti": [],
                                              "partecipanti": []})
    verifica("pptx: elenco vuoto, diapositiva tolta", len(Presentation(io.BytesIO(data)).slides)
             == 2)
    return cartella


# ─────────────────────────── servizio a voce ───────────────────────────

class Prof:
    def __init__(self, pid, name):
        self.id, self.name = pid, name


class Speakers:
    p = {"Dario": Prof("dario-id", "Dario"), "Bianca": Prof("bianca-id", "Bianca")}

    def get(self, n):
        return self.p.get(n)


class SpeakerCtx:
    def __init__(self, name, level):
        self.current_speaker, self.current_level, self.identified_by = name, level, "voce"


class Estrattore:
    """Le risposte dell'estrazione, preparate dalla prova una per chiamata."""

    def __init__(self):
        self.code, self.chiamate = [], []

    def __call__(self, messages, schema):
        self.chiamate.append((messages, schema))
        return json.dumps(self.code.pop(0) if self.code else {})


class Scrittore:
    def __init__(self):
        self.chiamate = []

    def __call__(self, sistema, utente):
        self.chiamate.append(utente)
        return "Il lavoro comprende la grafica e lo sviluppo del sito in cinque pagine."


def ufficio(nome, cartella_modelli=None, **cfgkw):
    cfg = Config()
    cfg.documenti_attesa_s = 5.0
    cfg.fatture_emittente = dict(EMITTENTE)
    for k, v in cfgkw.items():
        setattr(cfg, k, v)
    db = str(TMP / f"{nome}.db")
    docs = Documenti(cfg, db, writer=object(), delivery=LocalDelivery(TMP / nome),
                     formats=("word", "excel", "pdf"))
    est, scr = Estrattore(), Scrittore()
    u = Ufficio(cfg, db, docs, estrai=est, scrittore=scr, cartella_modelli=cartella_modelli,
                oggi=lambda: OGGI)
    u.rubrica.aggiungi(ROSSI, "casa", "dario-id")
    u.rubrica.aggiungi(PRIVATO, "casa", "dario-id")
    return cfg, u, est, scr


def chiama(reg, cfg, u, nome, args, chi="Dario", livello="amministra", frase="", turno=1):
    ctx = ToolContext(cfg=cfg, speakers=Speakers(), speaker_ctx=SpeakerCtx(chi, livello),
                      speaker=None, documenti=u.documenti, ufficio=u, user_text=frase,
                      turno=turno)
    out = json.loads(reg.call(nome, args, ctx, livello))
    out["_regole"] = list(ctx.regole)
    return out


def prova_servizio(cartella_modelli):
    cfg, u, est, scr = ufficio("voce", cartella_modelli)
    reg = build_registry(documenti=("word", "excel", "pdf"), ufficio=u.nomi_modelli())
    # 1. Mancano i dati: domanda con l'azione in sospeso
    est.code.append({"cliente": None, "righe": None})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "fammi una fattura"},
               frase="Calliope, fammi una fattura", turno=1)
    verifica("dati mancanti: una domanda, niente file", "mi servono" in r["risposta_finale"]
             and r["in_sospeso"]["argomenti"] == {"modello": "fattura"}, r["risposta_finale"])
    # 2. La risposta completa la bozza → proposta
    est.code.append({"cliente": "Rossi e Figli", "righe": [
        {"descrizione": "Consulenza sito", "quantita": 2, "prezzo": 500, "aliquota": None}]})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "a Rossi e figli, "
               "due giornate di consulenza a 500"}, turno=2)
    oid = r.get("proposta")
    verifica("proposta con il numero e i conti del programma",
             r["risposta_finale"].startswith("La fattura numero 1 del 2026 a Rossi & Figli")
             and "totale 1.220 euro" in r["risposta_finale"] and "da pagare 1.020" in
             r["risposta_finale"] and r["risposta_finale"].endswith("La preparo?")
             and r["in_sospeso"]["argomenti"]["proposta"] == oid, r["risposta_finale"])
    verifica("l'estrazione riceve i dati già raccolti", "Dati già raccolti" not in
             est.chiamate[0][0][1]["content"] and len(est.chiamate) == 2)
    # «Sì» nel turno sbagliato (dal 04/10 la proposta vale 3 turni: qui 4 dopo) o da
    # un'altra persona: niente
    r2 = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "sì",
                                                 "proposta": oid}, turno=6)
    r3 = chiama(reg, cfg, u, "modello_compila", {"modello": "preventivo", "dati": "x",
                                                 "proposta": "None"}, turno=3)
    verifica("proposta che non è un id ignorata (regola)", "ufficio_proposta_non_id" in
             r3["_regole"])
    verifica("«sì» fuori turno rifiutato", r2["ok"] is False and "ufficio_senza_offerta" in
             r2["_regole"])
    est.code.append({"cliente": "Rossi e figli", "righe": [{"descrizione": "Consulenza", "quantita": 2,
                                                    "prezzo": 500, "aliquota": None}]})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "uguale"}, turno=6)
    oid = r["proposta"]
    r2 = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "sì",
                                                 "proposta": oid}, chi="Bianca", turno=7)
    verifica("«sì» di un'altra persona rifiutato", r2["ok"] is False)
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "sì",
                                                "proposta": oid}, turno=7)
    verifica("«sì» nel turno dopo: fattura emessa, PDF e XML", r.get("ok") and r.get("numero")
             == 1 and r.get("xml", "").startswith("IT01234567897_") and "La apro" not in
             r["risposta_finale"] and "non lo invio" in r["risposta_finale"], r["risposta_finale"])
    cart = TMP / "voce" / "Fatture"
    pdf = next(cart.glob("*.pdf"))
    xml = (cart / r["xml"]).read_bytes()
    testo = plain_text("pdf", pdf.read_bytes())
    if testo:
        verifica("PDF di cortesia: numero, cliente, totali, IBAN, avviso",
                 all(s in testo for s in ("Fattura n. 1/2026", "Rossi & Figli", "1.220,00 €",
                                          "1.020,00 €", "Ritenuta", "IT60X0542811101000000123456",
                                          "Copia di cortesia")), testo[:150].replace("\n", " / "))
    verifica("XML: numero 1/2026, data, ritenuta", campo(xml, "Numero") == ["1/2026"]
             and campo(xml, "Data")[0] == OGGI.isoformat() and campo(xml, "ImportoRitenuta")
             == ["200.00"])
    if fatturapa.xsd_presente(XSD):
        verifica("XML emesso valido per l'XSD", fatturapa.valida_xsd(xml, XSD) == [])
    verifica("registro della numerazione", [d["numero"] for d in u.numeri.elenco("fatture", 2026)]
             == [1])
    # Privato: niente ritenuta (non ha partita IVA), bollo solo se serve
    est.code.append({"cliente": "Mario Rossi", "righe": [{"descrizione": "Riparazione PC",
                                                          "quantita": None, "prezzo": 122,
                                                          "aliquota": None}],
                     "iva_inclusa": True})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "a Mario Rossi "
               "riparazione 122 euro iva inclusa"}, turno=10)
    verifica("privato, IVA inclusa: totale 122, niente ritenuta, numero 2",
             "numero 2 del 2026" in r["risposta_finale"] and "totale 122 euro" in
             r["risposta_finale"] and "ritenuta" not in r["risposta_finale"], r["risposta_finale"])
    # Cliente che non c'è in rubrica
    est.code.append({"cliente": "Verdi", "righe": [{"descrizione": "x", "prezzo": 1,
                                                    "quantita": None, "aliquota": None}]})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "a Verdi"}, turno=12)
    verifica("cliente non in rubrica: lo dice, niente proposta", "Non trovo «Verdi»" in
             r["risposta_finale"] and "proposta" not in r)
    # Permessi: un familiare non fattura (predefinito), un ospite non vede il tool
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "fattura", "dati": "x"},
               chi="Bianca", livello="familiare", turno=14)
    verifica("familiare: le fatture solo a chi amministra", r["ok"] is False and
             "solo chi amministra" in r["risposta_finale"])
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "preventivo", "dati": "x"},
               chi="Bianca", livello="ospite", turno=15)
    verifica("ospite rifiutato dal registro", r.get("ok") is False and "NON" in r.get("fatto", ""))
    names_ospite = {s["function"]["name"] for s in reg.schemas_for("ospite")}
    verifica("l'ospite non può usare i tool dell'ufficio (permessi)", not names_ospite & {
        "modello_compila", "anagrafica_cerca", "anagrafica_salva"})
    # Preventivo con il testo lungo dello scrittore
    est.code.append({"cliente": "Rossi e figli", "oggetto": "Sito vetrina", "descrizione":
                     "grafica e sviluppo, cinque pagine", "righe": [
                         {"descrizione": "Grafica", "quantita": None, "prezzo": 500,
                          "aliquota": None}, {"descrizione": "Sviluppo", "quantita": None,
                                              "prezzo": 1000, "aliquota": None}],
                     "validita_giorni": 60})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "preventivo", "dati": "..."},
               chi="Bianca", livello="familiare", turno=20)
    oid = r.get("proposta")
    verifica("preventivo: proposta (niente ritenuta nei preventivi)", oid and "Il preventivo "
             "numero 1 del 2026" in r["risposta_finale"] and "ritenuta" not in
             r["risposta_finale"] and r["risposta_finale"].endswith("Lo preparo?"),
             r["risposta_finale"])
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "preventivo", "dati": "sì",
                                                "proposta": oid}, chi="Bianca",
               livello="familiare", turno=21)
    pdf = next((TMP / "voce" / "Preventivi").glob("*.pdf"))
    testo = plain_text("pdf", pdf.read_bytes())
    verifica("preventivo emesso: testo dello scrittore, validità, numero P1/2026",
             r.get("ok") and scr.chiamate and (not testo or (
                 "cinque pagine" in testo and "60 giorni" in testo and "P1/2026" in testo)),
             r["risposta_finale"])
    # DDT
    est.code.append({"cliente": "Rossi e figli", "righe": [{"descrizione": "Monitor 27 pollici",
                                                    "quantita": 2, "unita": "pz"}],
                     "colli": 1, "vettore": "BRT"})
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "ddt", "dati": "..."}, turno=30)
    r = chiama(reg, cfg, u, "modello_compila", {"modello": "ddt", "dati": "sì",
                                                "proposta": r["proposta"]}, turno=31)
    pdf = next((TMP / "voce" / "DDT").glob("*.pdf"))
    testo = plain_text("pdf", pdf.read_bytes())
    verifica("DDT emesso con numero, beni e vettore", r.get("ok") and (not testo or all(
        s in testo for s in ("DDT1/2026", "Monitor 27 pollici", "BRT", "Destinatario"))),
             r["risposta_finale"])
    if cartella_modelli:
        # Modello Word dell'utente con serie: proposta, poi il file
        est.code.append({"cliente": "Rossi e figli", "oggetto": "Sito", "descrizione": "cinque pagine",
                         "voci": [{"descrizione": "Grafica", "quantita": None, "prezzo": 500,
                                   "aliquota": None}]})
        r = chiama(reg, cfg, u, "modello_compila", {"modello": "offerta sito", "dati": "..."},
                   turno=40)
        verifica("modello Word con serie: proposta", "Offerta per un sito numero 2 del 2026" in
                 r["risposta_finale"], r["risposta_finale"])
        r = chiama(reg, cfg, u, "modello_compila", {"modello": "offerta sito", "dati": "sì",
                                                    "proposta": r["proposta"]}, turno=41)
        f = TMP / "voce" / r.get("nome_file", "?")
        txt = plain_text("word", f.read_bytes()) if f.exists() else ""
        verifica("modello Word compilato: numero, cliente, voce, totale con IVA, testo lungo",
                 all(s in txt for s in ("Preventivo P2/2026", "Rossi & Figli", "Grafica | "
                                        "500,00 €", "610,00 €", "cinque pagine")),
                 txt[:160].replace("\n", " / "))
        est.code.append({"titolo": "di ottobre", "punti": ["Budget"], "partecipanti": None})
        r = chiama(reg, cfg, u, "modello_compila", {"modello": "verbale", "dati": "..."},
                   turno=50)
        verifica("modello PowerPoint: campo mancante → domanda", "partecipanti" in
                 r["risposta_finale"] and r.get("in_sospeso"), r["risposta_finale"])
        est.code.append({"titolo": "di ottobre", "punti": ["Budget"],
                         "partecipanti": ["Dario"]})
        r = chiama(reg, cfg, u, "modello_compila", {"modello": "verbale", "dati": "c'era Dario"},
                   turno=51)
        verifica("modello PowerPoint senza serie: fatto subito", r.get("ok") and
                 r.get("nome_file", "").endswith(".pptx"), r["risposta_finale"])
    # Rubrica a voce
    est.code.append({"denominazione": "Bianchi Snc", "partita_iva": "1234567890",
                     "comune": "Pavia"})
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "aggiungi", "nome": "Bianchi",
                                                 "dati": "partita IVA 1234567890"}, turno=60)
    verifica("rubrica: partita IVA sbagliata → la ridice", "non torna" in r["risposta_finale"]
             and "proposta" not in r, r["risposta_finale"])
    est.code.append({"denominazione": "Bianchi Snc", "partita_iva": "0 1 2 3 4 5 6 7 8 9 7",
                     "indirizzo": "Corso Cavour", "civico": "10", "cap": "27100",
                     "comune": "Pavia", "provincia": "PV", "solo_per_me": None})
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "aggiungi", "nome": "Bianchi",
                                                 "dati": "..."}, turno=61)
    verifica("rubrica: proposta con la partita IVA detta cifra per cifra",
             r["risposta_finale"].startswith("Aggiungo alla rubrica Bianchi Snc, cliente, "
                                             "partita IVA 0 1 2 3") and r.get("proposta"),
             r["risposta_finale"])
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "aggiungi", "nome": "Bianchi",
                                                 "proposta": r["proposta"]}, turno=62)
    r = chiama(reg, cfg, u, "anagrafica_cerca", {"testo": "Bianchi"}, turno=63)
    verifica("rubrica: aggiunto e ritrovato", "Bianchi Snc" in r["risposta_finale"]
             and "Pavia" in r["risposta_finale"], r["risposta_finale"])
    est.code.append({"email": "info@bianchi.it"})
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "modifica", "nome": "Bianchi",
                                                 "dati": "email info@bianchi.it"}, turno=64)
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "modifica", "nome": "Bianchi",
                                                 "proposta": r.get("proposta", "")}, turno=65)
    c, _ = u.rubrica.trova("Bianchi", "dario-id")
    verifica("rubrica: modifica confermata", c and c.get("email") == "info@bianchi.it")
    r = chiama(reg, cfg, u, "anagrafica_salva", {"azione": "elimina", "nome": "Bianchi"},
               turno=66)
    verifica("rubrica: togliere chiede conferma", r["risposta_finale"] == "Tolgo Bianchi Snc "
             "dalla rubrica?" and u.rubrica.trova("Bianchi", "dario-id")[0] is not None)
    # Emittente mancante
    cfg2, u2, est2, _ = ufficio("senza-emittente", fatture_emittente={})
    reg2 = build_registry(documenti=("pdf",), ufficio=u2.nomi_modelli())
    est2.code.append({"cliente": "Rossi", "righe": [{"descrizione": "x", "prezzo": 10,
                                                     "quantita": None, "aliquota": None}]})
    r = chiama(reg2, cfg2, u2, "modello_compila", {"modello": "fattura", "dati": "..."}, turno=1)
    verifica("senza i dati dell'emittente: dice dove metterli", "fatture_emittente" in
             r["risposta_finale"], r["risposta_finale"])
    # Prompt e prefisso
    base = build_registry(documenti=("word", "excel", "pdf"))
    s0 = len(json.dumps(base.all_schemas(), ensure_ascii=False))
    s1 = len(json.dumps(reg.all_schemas(), ensure_ascii=False))
    p0 = cfg.prompt_for(False, documenti=True)
    p1 = cfg.prompt_for(False, documenti=True, ufficio=True)
    verifica("prompt: nomina modello_compila e la rubrica solo con l'ufficio",
             "modello_compila" in p1 and "anagrafica_salva" in p1 and "modello_compila" not in p0)
    print(f"   prefisso: schemi dei tool +{s1 - s0} caratteri ({s0} → {s1}), prompt "
          f"+{len(p1) - len(p0)} caratteri")


if __name__ == "__main__":
    prova_codici()
    prova_conti()
    prova_xml()
    prova_numerazione()
    prova_rubrica()
    cart = prova_modelli()
    prova_servizio(cart)
    print(f"\n{'Tutto bene' if not errori else f'{errori} errori'}")
    sys.exit(1 if errori else 0)
