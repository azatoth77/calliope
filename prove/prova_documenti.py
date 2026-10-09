import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

"""Prova a secco dei documenti (calliope/documenti/, calliope/tools/documenti.py).

Niente Ollama: lo scrittore è finto e restituisce JSON preparati. Tutto in una cartella
temporanea, mai nella vera Documenti. Si provano i tre renderer (riletti con le stesse
librerie: accenti, «€», virgolette, tabella, formule), la validazione (tipi sbagliati,
blocchi sconosciuti, testo troppo lungo), le formule ammesse e la formula injection, i
nomi dei file ripuliti e senza sovrascrittura, la modifica che rigenera il file (e che non
sovrascrive un file cambiato a mano o aperto), il lavoro in secondo piano con l'annuncio,
i permessi (ospite rifiutato, documento di un'altra persona intoccabile), «aprilo» con il
PC finto, la risposta finale nel Brain e il prompt.
"""

import io
import json
import re
import tempfile
import threading
import time
from pathlib import Path

from calliope.brain import Brain
from calliope.config import Config, example_yaml, load_config
from calliope.documenti import Documenti, DocumentoNonValido, validate
from calliope.documenti import consegna as consegna_mod
from calliope.documenti.consegna import LocalDelivery, documents_folder, unique_path
from calliope.documenti.formato import (check_formula, classify_cell, describe_change,
                                        parse_number, safe_filename, summary)
from calliope.documenti.render import find_font, plain_text, render, sheet_totals
from calliope.documenti.scrittore import Writer
from calliope.tools.builtin import build_registry
from calliope.tools.spec import ToolContext
from prove.pc_finto import FakePC

errori = 0


def verifica(nome, ok, dettaglio=""):
    global errori
    errori += not ok
    print(f"{'ok ' if ok else 'ERR'} {nome}" + (f"  {dettaglio}" if dettaglio else ""))


def errore_di(formato, data) -> list[str]:
    try:
        validate(formato, data)
        return []
    except DocumentoNonValido as e:
        return e.errors


TMP = Path(tempfile.mkdtemp(prefix="calliope-documenti-"))

LETTERA = {"titolo": "Disdetta palestra: “Fit & Go”", "blocchi": [
    {"tipo": "paragrafo", "testo": "Roma, 27 settembre 2026", "allinea": "destra"},
    {"tipo": "paragrafo", "testo": "Spett.le Palestra Fit & Go\n[indirizzo della palestra]"},
    {"tipo": "paragrafo", "testo": "Oggetto: disdetta dell’abbonamento\n\nCon la presente "
     "comunico la disdetta perché mi trasferirò. La quota di 45 € al mese non sarà più dovuta "
     "già da ottobre: è così."},
    {"tipo": "paragrafo", "testo": "Cordiali saluti,\nDario", "allinea": "destra"}]}
SPESE_TESTO = {"titolo": "Spese di settembre", "blocchi": [
    {"tipo": "titolo", "testo": "Spese di settembre"},
    {"tipo": "tabella", "colonne": ["Voce", "Importo (€)"],
     "righe": [["Affitto", "800"], ["Luce", "90"], ["Gas", "60"]], "totale": True},
    {"tipo": "elenco", "voci": ["- pagare entro il 5", "controllare la bolletta"]}]}
SPESE_FOGLIO = {"titolo": "Spese di settembre", "fogli": [
    {"nome": "Settembre", "colonne": ["Voce", "Importo (€)", "Doppio"],
     "righe": [["Affitto", "800", "=B2*2"], ["Luce", "90,50", "=SOMMA(B3;B3)"],
               ["Gas", 60, "=ROUND(B4*2,0)"], ["+39 333 1234567", "", ""],
               ["@nota", "-5", ""], ["- da pagare", "", ""]],
     "totale": True}]}


# ─────────────────────────── renderer ───────────────────────────

def prova_render():
    doc = validate("word", LETTERA)
    verifica("paragrafo con righe vuote diviso in più paragrafi",
             [b["testo"][:8] for b in doc["blocchi"]]
             == ["Roma, 27", "Spett.le", "Oggetto:", "Con la p", "Cordiali"])
    data = render("word", doc)
    text = plain_text("word", data)
    verifica("docx: accenti, € e virgolette", all(s in text for s in
                                                  ("perché", "45 €", "dell’abbonamento", "è così")),
             repr(text[:80]))
    from docx import Document
    d = Document(io.BytesIO(data))
    verifica("docx: titolo nelle proprietà", d.core_properties.title == LETTERA["titolo"])
    verifica("docx: data allineata a destra", d.paragraphs[0].alignment is not None)
    verifica("docx: a capo dentro il paragrafo", "Palestra Fit & Go\n[indirizzo" in text)

    data = render("word", validate("word", SPESE_TESTO))
    d = Document(io.BytesIO(data))
    righe = [[c.text for c in r.cells] for r in d.tables[0].rows]
    verifica("docx: tabella con il totale calcolato",
             righe[-1] == ["Totale", "950 €"] and righe[0] == ["Voce", "Importo (€)"], str(righe))
    verifica("docx: elenco senza il trattino del modello",
             any(p.text == "pagare entro il 5" for p in d.paragraphs))

    font = find_font(Config().documenti_font)
    verifica("font TrueType di sistema trovato", font is not None, str(font))
    data = render("pdf", validate("pdf", LETTERA), Config().documenti_font)
    verifica("pdf: file PDF", data[:5] == b"%PDF-")
    text = plain_text("pdf", data)
    if text:
        verifica("pdf: accenti, € e virgolette", all(s in text for s in ("perché", "€", "’", "è così")),
                 repr(text[:120]))
    else:
        print("   (pypdf non installato: testo del PDF non riletto)")
    data = render("pdf", validate("pdf", SPESE_TESTO), Config().documenti_font)
    text = plain_text("pdf", data)
    if text:
        verifica("pdf: tabella e totale", "Affitto" in text and "950 €" in text, repr(text[:200]))
    # Titolo del documento in cima se il contenuto non ne ha uno (01/10: «metti in alto il
    # titolo» cambiava solo le proprietà del file); nelle lettere no
    compiti = {"titolo": "Compiti Matteo", "blocchi": [
        {"tipo": "elenco", "voci": ["Matematica: pagina 45", "Inglese: esercizio 3"]}]}
    text = plain_text("pdf", render("pdf", validate("pdf", compiti), Config().documenti_font))
    if text:
        verifica("pdf: titolo del documento in cima", text.lstrip().startswith("Compiti Matteo"),
                 repr(text[:80]))
    d = Document(io.BytesIO(render("word", validate("word", compiti))))
    verifica("word: titolo del documento in cima", d.paragraphs[0].text == "Compiti Matteo")
    lettera = {"titolo": "Disdetta palestra", "blocchi": [{"tipo": "paragrafo", "testo": "Gentili signori,"}]}
    d = Document(io.BytesIO(render("word", lettera)))
    verifica("word: lettera senza intestazione in cima", d.paragraphs[0].text == "Gentili signori,")
    # Senza font TrueType: Helvetica, «€» → «EUR», le lettere accentate restano
    data = render("pdf", validate("pdf", LETTERA), [])
    text = plain_text("pdf", data)
    if text:
        verifica("pdf senza TrueType: ripiego leggibile", "EUR" in text and "perché" in text,
                 repr(text[:120]))

    doc = validate("excel", SPESE_FOGLIO)
    data = render("excel", doc)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data))["Settembre"]
    verifica("xlsx: numeri veri", ws["B2"].value == 800 and ws["B3"].value == 90.5
             and ws["B4"].value == 60 and ws["B6"].value == -5,
             f"{ws['B2'].value!r} {ws['B3'].value!r} {ws['B4'].value!r}")
    verifica("xlsx: formule ammesse (anche all'italiana)",
             (ws["C2"].value, ws["C3"].value, ws["C4"].value)
             == ("=B2*2", "=SUM(B3,B3)", "=ROUND(B4*2,0)"),
             f"{ws['C2'].value} {ws['C3'].value} {ws['C4'].value}")
    verifica("xlsx: riga del totale con SUM", (ws["A8"].value, ws["B8"].value)
             == ("Totale", "=SUM(B2:B7)"), f"{ws['A8'].value} {ws['B8'].value}")
    verifica("xlsx: formato in euro", "€" in ws["B2"].number_format and "€" in ws["B8"].number_format)
    verifica("xlsx: testo con + e @ scritto come testo protetto",
             ws["A5"].data_type == "s" and ws["A5"].quotePrefix and ws["A6"].data_type == "s"
             and ws["A6"].quotePrefix and ws["A7"].value == "- da pagare" and ws["A7"].quotePrefix)
    verifica("xlsx: intestazione in grassetto e bloccata", ws["A1"].font.b and ws.freeze_panes == "A2")
    verifica("totale del foglio calcolato in Python", sheet_totals(doc) == [("Importo (€)", "945,50 €")],
             str(sheet_totals(doc)))


# ─────────────────────────── validazione e formule ───────────────────────────

def prova_validazione():
    verifica("tipo di blocco sconosciuto", any("sconosciuto" in e for e in errore_di(
        "word", {"titolo": "x", "blocchi": [{"tipo": "immagine", "url": "http://x"}]})))
    verifica("blocchi non è un elenco", errore_di("word", {"titolo": "x", "blocchi": "ciao"}) != [])
    verifica("testo di tipo sbagliato", any("testo" in e for e in errore_di(
        "word", {"titolo": "x", "blocchi": [{"tipo": "paragrafo", "testo": {"a": 1}}]})))
    verifica("testo troppo lungo", any("troppo lungo" in e for e in errore_di(
        "word", {"titolo": "x", "blocchi": [{"tipo": "paragrafo", "testo": "a" * 3500}]})))
    verifica("titolo mancante", any("titolo" in e for e in errore_di(
        "pdf", {"blocchi": [{"tipo": "paragrafo", "testo": "ciao"}]})))
    verifica("riga con più valori delle colonne", any("colonne" in e for e in errore_di(
        "word", {"titolo": "x", "blocchi": [{"tipo": "tabella", "colonne": ["a"],
                                              "righe": [["1", "2"]]}]})))
    verifica("troppe colonne", errore_di("word", {"titolo": "x", "blocchi": [
        {"tipo": "tabella", "colonne": [str(i) for i in range(20)], "righe": []}]}) != [])
    ok = validate("word", {"titolo": " x ", "blocchi": [{"tipo": "tabella", "colonne": ["a", "b"],
                                                          "righe": [["1"], [2, 3, ""]]}]})
    verifica("righe corte completate, numeri in testo",
             ok["blocchi"][0]["righe"] == [["1", ""], ["2", "3"]] and ok["titolo"] == "x",
             str(ok))
    verifica("non JSON", any("JSON" in e for e in errore_di("word", "{non json")))
    verifica("foglio senza colonne", errore_di("excel", {"titolo": "x", "fogli": [
        {"nome": "a", "righe": []}]}) != [])
    doc = validate("excel", {"titolo": "x", "fogli": [
        {"nome": "Spese: [ott]", "colonne": ["a"], "righe": []},
        {"nome": "spese  ott ", "colonne": ["a"], "righe": []}]})
    verifica("nomi dei fogli ripuliti e unici", [s["nome"] for s in doc["fogli"]]
             == ["Spese ott", "spese ott 2"], str([s["nome"] for s in doc["fogli"]]))

    # Formule: elenco consentito e riferimenti dentro la tabella
    for f in ("=HYPERLINK(\"http://x\",\"clic\")", "=WEBSERVICE(\"http://x\")",
              "=cmd|' /C calc'!A0", "=IMPORTXML(A1)", "=[1]Foglio1!A1", "=Foglio2!A1",
              "=INDIRECT(\"A1\")", "=SUM(A1:A5", "=B99", "=SUM(Z1)", "=A1+"):
        try:
            check_formula(f, 3, 5)
            bad = False
        except ValueError:
            bad = True
        # «=A1+» passa il tokenizzatore: lo scarta Excel, ma non è pericolosa
        if f == "=A1+":
            continue
        verifica(f"formula rifiutata: {f}", bad)
    verifica("formula con IF e stringhe", check_formula('=IF(B2>100,"alto","basso")', 3, 5)
             == '=IF(B2>100,"alto","basso")')
    verifica("formula all'italiana con ;", check_formula("=media(b2:b4)", 3, 5) == "=AVERAGE(B2:B4)")
    verifica("cella «=HYPERLINK» = errore di validazione", any("HYPERLINK" in e for e in errore_di(
        "excel", {"titolo": "x", "fogli": [{"nome": "a", "colonne": ["a", "b"],
                                             "righe": [["x", '=HYPERLINK("http://x")']]}]})))
    verifica("cella «@» = testo protetto", classify_cell("@SUM(A1)", 2, 3) == ("testo", "@SUM(A1)", True))
    verifica("cella «-5» = numero", classify_cell("-5", 2, 3)[:2] == ("numero", -5))
    verifica("numeri all'italiana", [parse_number(x) for x in ("1.200,50", "€ 90", "12.5", "007", "15/10")]
             == [(1200.5, False), (90, True), (12.5, False), None, None])


# ─────────────────────────── nomi dei file ───────────────────────────

def prova_nomi():
    verifica("nome ripulito", safe_filename('Lettera: disdetta/palestra? "Fit" <ok>.')
             == "Lettera disdetta palestra Fit ok", safe_filename('Lettera: disdetta/palestra? "Fit" <ok>.'))
    verifica("nome riservato", safe_filename("con") == "con_")
    verifica("nome vuoto", safe_filename(" ?? ") == "Documento")
    verifica("nome lungo tagliato", len(safe_filename("a" * 300)) == 80)
    folder = TMP / "nomi"
    folder.mkdir()
    (folder / "Spese.xlsx").write_bytes(b"x")
    (folder / "Spese (2).xlsx").write_bytes(b"x")
    verifica("senza sovrascrivere", unique_path(folder, "Spese", "xlsx").name == "Spese (3).xlsx")
    d = LocalDelivery(folder)
    r = d.deliver("Spese", "xlsx", b"nuovo")
    verifica("consegna: nome libero", r["nome_file"] == "Spese (3).xlsx"
             and (folder / "Spese.xlsx").read_bytes() == b"x")
    verifica("consegna: dove, a voce", d.where() == "nella cartella nomi"
             and LocalDelivery().where() == "nella cartella Calliope dei Documenti")
    verifica("cartella Documenti di Windows", documents_folder().is_dir(), str(documents_folder()))
    verifica("nessun file temporaneo rimasto", not list(folder.glob(".calliope-*")))


# ─────────────────────────── servizio e tool ───────────────────────────

class FakeWriter:
    """Scrittore finto: documenti preparati e modifiche a regole, con un ritardo."""

    def __init__(self, delay=0.0):
        self.delay = delay
        self.calls = []
        self.last_stats = {"s": 0.0, "token": 0, "tentativi": 1}

    def write(self, formato, richiesta, detto="", persona=None, titolo=""):
        self.calls.append(("write", formato, richiesta))
        time.sleep(self.delay)
        if formato == "excel":
            return validate("excel", {"titolo": "Spese di settembre", "fogli": [
                {"nome": "Spese", "colonne": ["Voce", "Importo (€)"],
                 "righe": [["Affitto", "800"], ["Luce", "90"], ["Gas", "60"]], "totale": True}]})
        if "vuoto" in richiesta:
            raise DocumentoNonValido(["blocchi: serve un elenco di blocchi non vuoto"])
        return validate(formato, LETTERA)

    def edit(self, formato, doc, modifica, detto=""):
        self.calls.append(("edit", formato, modifica))
        time.sleep(self.delay)
        new = json.loads(json.dumps(doc))
        if "riga" in modifica:
            new["fogli"][0]["righe"].append(["Internet", "30"])
        elif "data" in modifica:
            day = re.search(r"\d+ ottobre", modifica).group()
            new["blocchi"][0]["testo"] = f"Roma, {day} 2026"
        return validate(formato, new)


class Prof:
    def __init__(self, pid, name, admin=False):
        self.id, self.name, self.admin = pid, name, admin


class Speakers:
    def __init__(self):
        self.p = {"Dario": Prof("dario-id", "Dario", True), "Bianca": Prof("bianca-id", "Bianca")}

    def get(self, n):
        return self.p.get(n)

    def known_speakers(self):
        return list(self.p)


class SpeakerCtx:
    def __init__(self, name, level):
        self.current_speaker, self.current_level, self.from_session = name, level, False


def servizio(cartella, delay=0.0, pcs=None, on_done=None):
    cfg = Config()
    cfg.documenti_attesa_s = 2.0
    svc = Documenti(cfg, str(TMP / f"{cartella}.db"), writer=FakeWriter(delay),
                    delivery=LocalDelivery(TMP / cartella), on_done=on_done, pcs=pcs)
    return cfg, svc


def chiama(reg, svc, cfg, nome, args, chi, livello, frase="", pc=None):
    ctx = ToolContext(cfg=cfg, speakers=Speakers(), speaker_ctx=SpeakerCtx(chi, livello),
                      speaker=None, documenti=svc, pc=pc, user_text=frase)
    return json.loads(reg.call(nome, args, ctx, livello))


def prova_servizio():
    pc = FakePC()
    pcs = {"portatile": pc}
    cfg, svc = servizio("docs", pcs=pcs)
    reg = build_registry(pc=pcs, documenti=svc.formati)
    names = {s["function"]["name"] for s in reg.schemas_for("familiare")}
    verifica("tool registrati per i familiari", {"documento_crea", "documento_modifica"} <= names)
    verifica("niente tool per gli ospiti (permessi)", not {"documento_crea", "documento_modifica"}
             & {s["function"]["name"] for s in reg.schemas_for("ospite")})
    verifica("senza documenti niente tool", "documento_crea" not in
             {s["function"]["name"] for s in build_registry().all_schemas()})
    enum = next(s for s in reg.all_schemas() if s["function"]["name"] == "documento_crea")
    verifica("enum dei formati", enum["function"]["parameters"]["properties"]["formato"]["enum"]
             == ["word", "excel", "pdf"])

    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "excel", "richiesta": "spese"},
               None, "ospite", "fammi una tabella")
    verifica("ospite rifiutato con «NON»", r.get("ok") is False and "NON" in r.get("fatto", ""))
    verifica("ospite: nessun file", not (TMP / "docs").exists() or not list((TMP / "docs").iterdir()))

    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "excel", "richiesta":
                                                 "spese di settembre: affitto 800, luce 90, gas 60, con il totale"},
               "Bianca", "familiare", "fammi una tabella excel con le spese di settembre", pcs)
    path = TMP / "docs" / "Spese di settembre.xlsx"
    verifica("crea: file scritto", path.exists(), str(r))
    verifica("crea: risposta finale con cosa contiene e dove",
             r.get("risposta_finale") == "Ho preparato il foglio Excel «Spese di settembre»: una "
             "tabella di 3 righe con il totale, nella cartella docs. Lo apro?", r.get("risposta_finale"))
    verifica("crea: il documento non si legge", "Affitto" not in json.dumps(r, ensure_ascii=False))
    verifica("crea: «Lo apro?» diventa un'azione in sospeso (pc_apri_file, risultato 1)",
             r.get("in_sospeso") == {"domanda": "Lo apro?",
                                     "cosa": "il foglio Excel «Spese di settembre»",
                                     "tool": "pc_apri_file", "argomenti": {"risultato": 1}},
             str(r.get("in_sospeso")))
    last = svc.archive.last("bianca-id")
    verifica("crea: JSON archiviato a parte", last and last["contenuto"]["fogli"][0]["righe"][0]
             == ["Affitto", "800"] and last["rif"] == str(path))

    # «Aprilo»: Bianca non è proprietaria del PC, ma il documento è suo
    r = chiama(reg, svc, cfg, "pc_apri_file", {"risultato": 1}, "Bianca", "familiare", "aprilo", pcs)
    verifica("aprilo: pc_apri_file(1) apre il documento", pc.azioni[-1:] == [("apri", str(path))], str(r))
    r = chiama(reg, svc, cfg, "pc_cerca_file", {"testo": "bolletta"}, "Bianca", "familiare", pc=pcs)
    verifica("ma le ricerche di file restano del proprietario", r.get("ok") is False)

    # Stesso titolo: file nuovo, non sovrascritto
    chiama(reg, svc, cfg, "documento_crea", {"formato": "excel", "richiesta": "spese"},
           "Bianca", "familiare", "rifammi la tabella", pcs)
    verifica("stesso titolo: «(2)»", (TMP / "docs" / "Spese di settembre (2).xlsx").exists())

    # Modifica: stesso file, versione precedente conservata nel database
    before = path.stat().st_mtime
    r = chiama(reg, svc, cfg, "documento_modifica", {"modifica": "aggiungi una riga: internet 30"},
               "Bianca", "familiare", "aggiungi una riga: internet 30", pcs)
    last = svc.archive.last("bianca-id")
    verifica("modifica: riga in più, totale ricalcolato", r.get("risposta_finale")
             == "Fatto, ho aggiornato «Spese di settembre»: una riga in più.", r.get("risposta_finale"))
    target = TMP / "docs" / "Spese di settembre (2).xlsx"   # l'ultimo documento di Bianca
    from openpyxl import load_workbook
    ws = load_workbook(target).active
    verifica("modifica: file rigenerato con la SUM giusta", (ws["A5"].value, ws["B6"].value)
             == ("Internet", "=SUM(B2:B5)"), f"{ws['A5'].value} {ws['B6'].value}")
    verifica("modifica: versione 2 e versione 1 archiviata", last["versione"] == 2
             and len(svc.archive.versions(last["id"])) == 1)
    verifica("il primo file non è stato toccato", path.stat().st_mtime == before)

    # Il documento di un'altra persona non si modifica
    r = chiama(reg, svc, cfg, "documento_modifica", {"modifica": "aggiungi una riga: pane 3"},
               "Dario", "amministra", "aggiungi una riga", pcs)
    verifica("Dario non modifica il documento di Bianca", r.get("ok") is False
             and "Non trovo un tuo documento" in r.get("risposta_finale", ""), str(r))
    ws = load_workbook(target).active
    verifica("il file di Bianca è rimasto uguale", ws["A6"].value == "Totale")

    # Lettera di Dario, poi modifica della data
    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "word", "richiesta":
                                                 "lettera di disdetta della palestra"},
               "Dario", "amministra", "preparami una lettera di disdetta per la palestra", pcs)
    verifica("lettera: frase al femminile", r.get("risposta_finale", "").startswith(
        "Ho preparato la lettera «Disdetta palestra: “Fit & Go”»: 5 paragrafi,")
        and r["risposta_finale"].endswith("La apro?"), r.get("risposta_finale"))
    letter = TMP / "docs" / "Disdetta palestra Fit & Go.docx"
    verifica("lettera: nome del file ripulito", letter.exists(), str(list((TMP / "docs").iterdir())))

    # Cambiata a mano (data di modifica diversa): la nuova versione va in un file nuovo
    os.utime(letter, (time.time() + 100, time.time() + 100))
    r = chiama(reg, svc, cfg, "documento_modifica", {"modifica": "cambia la data in 15 ottobre"},
               "Dario", "amministra", "cambia la data in 15 ottobre", pcs)
    verifica("cambiato a mano: nuova versione accanto", "cambiato a mano" in r.get("risposta_finale", "")
             and (TMP / "docs" / "Disdetta palestra Fit & Go (2).docx").exists(), r.get("risposta_finale"))
    text = plain_text("word", (TMP / "docs" / "Disdetta palestra Fit & Go (2).docx").read_bytes())
    verifica("modifica della data nel file", "15 ottobre 2026" in text and "27 settembre" not in text)

    # Aperto in Word (file bloccato): os.replace fallisce, la nuova versione va accanto
    real = consegna_mod.os.replace

    def locked(src, dst):
        if str(dst).endswith("(2).docx"):
            raise PermissionError(13, "file in uso")
        return real(src, dst)
    consegna_mod.os.replace = locked
    try:
        r = chiama(reg, svc, cfg, "documento_modifica", {"modifica": "cambia la data in 16 ottobre"},
                   "Dario", "amministra", "cambia la data", pcs)
    finally:
        consegna_mod.os.replace = real
    verifica("file aperto: salvato accanto e detto", "era aperto" in r.get("risposta_finale", "")
             and (TMP / "docs" / "Disdetta palestra Fit & Go (3).docx").exists(), r.get("risposta_finale"))
    verifica("nessun file temporaneo rimasto", not list((TMP / "docs").glob(".calliope-*")))

    # JSON invalido anche al secondo tentativo: errore chiaro, niente file
    n = len(list((TMP / "docs").iterdir()))
    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "pdf", "richiesta": "un documento vuoto"},
               "Dario", "amministra", "fai un pdf vuoto", pcs)
    verifica("JSON invalido: errore detto, nessun file", r.get("ok") is False
             and "Non sono riuscita" in r.get("risposta_finale", "")
             and len(list((TMP / "docs").iterdir())) == n, r.get("risposta_finale"))

    # Il formato detto a voce vince su quello del modello
    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "word", "richiesta": "lettera di disdetta"},
               "Dario", "amministra", "fammi un PDF con la lettera di disdetta", pcs)
    verifica("formato detto a voce: pdf", r.get("formato") == "pdf", str(r.get("formato")))
    svc.close()


def prova_secondo_piano():
    fired = threading.Event()
    cfg, svc = servizio("lenti", delay=0.3, on_done=fired.set)
    cfg.documenti_attesa_s = 0.05
    reg = build_registry(documenti=svc.formati)
    t0 = time.perf_counter()
    r = chiama(reg, svc, cfg, "documento_crea", {"formato": "word", "richiesta": "lettera di disdetta"},
               "Dario", "amministra", "preparami una lettera di disdetta")
    dt = time.perf_counter() - t0
    verifica("lento: risposta subito, in preparazione", r.get("in_preparazione") and dt < 0.25
             and r.get("risposta_finale") == "Te la preparo: ci vuole qualche secondo, ti avviso "
             "quando è pronta.", f"{dt:.2f}s {r.get('risposta_finale')}")
    # Modifica chiesta prima che la creazione finisca: in coda, legge il documento nuovo
    r = chiama(reg, svc, cfg, "documento_modifica", {"modifica": "cambia la data in 15 ottobre"},
               "Dario", "amministra", "cambia la data")
    verifica("modifica in coda dietro la creazione", r.get("in_preparazione") is True, str(r))
    verifica("a lavoro finito si sveglia l'ascolto", fired.wait(3))
    t0 = time.perf_counter()
    while svc.busy() and time.perf_counter() - t0 < 3:
        time.sleep(0.02)
    msgs = []
    while not svc.done.empty():
        msgs.append(svc.done.get_nowait()["messaggio"])
    verifica("annunci: pronta, poi aggiornata", len(msgs) == 2
             and msgs[0].startswith("Dario, è pronta la lettera «Disdetta palestra")
             and msgs[1].startswith("Dario, ho aggiornato «Disdetta palestra"), str(msgs))
    verifica("annuncio senza «lo apro» se il PC non c'è", "apro" not in msgs[0])
    svc.close()


class SulServer(LocalDelivery):
    """Consegna di Calliope sulla DGX con un telefono collegato: il satellite non riceve
    documenti e il file resta sul server (come consegna.RemoteDelivery._sul_server)."""

    def deliver(self, *a, **k):
        d = super().deliver(*a, **k)
        d["dove"] = "sul server, perché il satellite collegato non riceve documenti"
        d["remoto"] = False
        return d


class Schermi:
    """Schermi finti: `personale` = chi ha uno schermo personale aperto adesso."""

    def __init__(self, personale=("dario-id",)):
        self.personale = set(personale)
        self.inviate = []

    def mittente(self, ctx):
        from types import SimpleNamespace
        prof = ctx.speakers.get(ctx.speaker_ctx.current_speaker)
        return SimpleNamespace(persona=getattr(prof, "id", None))

    def personale_collegato(self, m):
        return m.persona in self.personale

    def invia(self, card, m, forza=False):
        if m.persona in self.personale:
            self.inviate.append((m.persona, card, forza))
            return {"schermi": ["telefono"], "destinatari": ["telefono"], "motivo": ""}
        return {"schermi": [], "destinatari": [], "motivo": "personale"}


def chiama_s(reg, svc, cfg, nome, args, chi, frase="", schermi=None):
    ctx = ToolContext(cfg=cfg, speakers=Speakers(), speaker_ctx=SpeakerCtx(chi, "amministra"),
                      speaker=None, documenti=svc, user_text=frase)
    ctx.schermi = schermi
    ctx.regole = []
    return json.loads(reg.call(nome, args, ctx, "amministra")), ctx


def prova_rilettura():
    """Casi veri della DGX del 09/10 (18:56–19:13), con nomi di fantasia: un foglio Excel creato
    dal telefono restava «sul server, perché il satellite non riceve documenti» e non si poteva
    più rileggere (lavoro_risultato ripiegava su un lavoro dell'agente)."""
    cfg = Config()
    cfg.documenti_attesa_s = 2.0
    svc = Documenti(cfg, str(TMP / "server.db"), writer=FakeWriter(),
                    delivery=SulServer(TMP / "server"))
    reg = build_registry(documenti=svc.formati, agenti=True)
    tel = Schermi()
    r, ctx = chiama_s(reg, svc, cfg, "documento_crea", {"formato": "excel", "richiesta": "spese"},
                      "Dario", "fammi una tabella delle spese", tel)
    f = r.get("risposta_finale") or ""
    verifica("sul server con il telefono aperto: la frase dice la scheda con «Scarica»",
             "sulla scheda del tuo schermo" in f and "«Scarica»" in f and "sul server" not in f
             and "documento_sulla_scheda" in ctx.regole, f)
    verifica("…e le frasi alternative non arrivano al modello",
             not {"frase_scheda", "annuncio_scheda", "frase", "annuncio"} & set(r), str(set(r)))
    verifica("cosa_fare nomina documento_leggi", "documento_leggi" in (r.get("cosa_fare") or ""))
    r, _ = chiama_s(reg, svc, cfg, "documento_crea", {"formato": "excel", "richiesta": "spese"},
                    "Dario", "fammi una tabella delle spese", Schermi(personale=()))
    verifica("contrario: senza uno schermo personale aperto, la frase di sempre (sul server)",
             "sul server, perché il satellite" in (r.get("risposta_finale") or ""),
             r.get("risposta_finale"))
    cfg_l, svc_l = servizio("portatile")
    reg_l = build_registry(documenti=svc_l.formati)
    r, _ = chiama_s(reg_l, svc_l, cfg_l, "documento_crea", {"formato": "excel",
                                                            "richiesta": "spese"},
                    "Dario", "fammi una tabella delle spese", tel)
    verifica("contrario: consegnato davvero (sul portatile) → la frase di sempre",
             "sulla scheda" not in (r.get("risposta_finale") or ""), r.get("risposta_finale"))
    # documento_leggi: «cosa c'è in quel file?», per titolo, l'ultimo, un altro
    names = {s["function"]["name"] for s in reg.schemas_for("familiare")}
    verifica("documento_leggi registrato per i familiari, non per gli ospiti",
             "documento_leggi" in names and "documento_leggi" not in
             {s["function"]["name"] for s in reg.schemas_for("ospite")})
    r, ctx = chiama_s(reg, svc, cfg, "documento_leggi", {}, "Dario",
                      "raccontami cosa c'è in quel file")
    verifica("documento_leggi: l'ultimo documento, con il contenuto per il modello",
             r.get("ok") and "Affitto" in (r.get("contenuto") or "")
             and "frasi" in (r.get("cosa_fare") or "") and "documento_letto" in ctx.regole,
             str(r)[:300])
    r, _ = chiama_s(reg, svc, cfg, "documento_leggi", {"documento": "il foglio Excel delle "
                                                                     "spese di settembre"},
                    "Dario", "riassumimi le spese")
    verifica("documento_leggi: per parole del titolo", r.get("ok") and r.get("titolo") ==
             "Spese di settembre", str(r)[:200])
    r, ctx = chiama_s(reg, svc, cfg, "documento_leggi", {"documento": "idratazione della pizza"},
                      "Dario", "e la pizza?")
    verifica("documento_leggi: un titolo che non c'è → errore con i documenti veri, mai un "
             "altro documento", r.get("ok") is False and r.get("correggibile") is True
             and "Spese di settembre" in (r.get("documenti") or []) and "contenuto" not in r
             and "documento_non_trovato" in ctx.regole, str(r))
    r, _ = chiama_s(reg, svc, cfg, "documento_leggi", {}, "Bianca", "cosa c'è nel file?")
    verifica("documento_leggi: i documenti di un'altra persona no", r.get("ok") is False
             and "contenuto" not in r, str(r))
    tel = Schermi()
    r, ctx = chiama_s(reg, svc, cfg, "documento_leggi", {"modo": "mostra"}, "Dario", "aprilo",
                      tel)
    verifica("documento_leggi mostra («aprilo»): sulla scheda del telefono, con «Scarica»",
             r.get("ok") and "sullo schermo" in (r.get("risposta_finale") or "")
             and tel.inviate and tel.inviate[-1][1].get("scarica"), str(r))
    r, _ = chiama_s(reg, svc, cfg, "documento_leggi", {"modo": "mostra"}, "Dario", "aprilo",
                    Schermi(personale=()))
    verifica("documento_leggi mostra senza schermo: il contenuto, da dire a voce",
             r.get("ok") and "contenuto" in r and "Non vedo un tuo schermo"
             in (r.get("cosa_fare") or ""), str(r)[:200])
    # In secondo piano: l'annuncio dice la scheda se la scheda è arrivata
    cfg2 = Config()
    cfg2.documenti_attesa_s = 0.05
    svc2 = Documenti(cfg2, str(TMP / "server2.db"), writer=FakeWriter(delay=0.3),
                     delivery=SulServer(TMP / "server2"))
    reg2 = build_registry(documenti=svc2.formati)
    r, _ = chiama_s(reg2, svc2, cfg2, "documento_crea", {"formato": "excel",
                                                         "richiesta": "spese"},
                    "Dario", "fammi una tabella", Schermi())
    t0 = time.perf_counter()
    while svc2.done.empty() and time.perf_counter() - t0 < 3:
        time.sleep(0.02)
    msg = svc2.done.get_nowait()["messaggio"] if not svc2.done.empty() else ""
    verifica("in secondo piano: l'annuncio dice la scheda con «Scarica»",
             r.get("in_preparazione") and "sulla scheda del tuo schermo" in msg, msg)
    for s in (svc, svc_l, svc2):
        s.close()


def prova_scrittore():
    """Un solo nuovo tentativo, con gli errori; poi errore chiaro."""
    w = Writer(Config())
    risposte = ['{"titolo":"x","blocchi":[{"tipo":"immagine"}]}',
                '{"titolo":"Compiti","blocchi":[{"tipo":"elenco","voci":["a","b"]}]}']
    visti = []

    def ask(messages, schema):
        visti.append(messages)
        w.last_stats = {"s": 0.1, "token": 10}
        return risposte.pop(0)
    w._ask = ask
    doc = w.write("pdf", "elenco dei compiti", "fai un pdf con i compiti", "Dario")
    verifica("scrittore: secondo tentativo riuscito", doc["blocchi"][0]["voci"] == ["a", "b"]
             and w.last_stats["tentativi"] == 2)
    verifica("scrittore: gli errori tornano al modello", "sconosciuto" in visti[1][-1]["content"])
    verifica("scrittore: data di oggi e persona nel prompt", "Oggi è" in visti[0][0]["content"]
             and "Dario" in visti[0][0]["content"])
    risposte[:] = ['{"titolo":"x","blocchi":[]}', '{"titolo":"x","blocchi":[]}']
    try:
        w.write("word", "x")
        bad = False
    except DocumentoNonValido:
        bad = True
    verifica("scrittore: due volte invalido → DocumentoNonValido", bad)
    # Modifica che restituisce solo un pezzo: rifiutata
    old = validate("word", LETTERA)
    risposte[:] = ['{"titolo":"x","blocchi":[{"tipo":"paragrafo","testo":"Roma, 15 ottobre"}]}',
                   json.dumps({**LETTERA, "blocchi": [{"tipo": "paragrafo", "testo": "Roma, 15 ottobre"}]
                               + LETTERA["blocchi"][1:]})]
    new = w.edit("word", old, "cambia la data in 15 ottobre")
    verifica("modifica parziale rifiutata, poi intera", len(new["blocchi"]) == 5
             and new["blocchi"][0]["testo"] == "Roma, 15 ottobre")


def prova_frasi():
    doc = validate("word", LETTERA)
    verifica("riassunto lettera", summary("word", doc, lettera=True) == "5 paragrafi")
    verifica("riassunto tabella", summary("word", validate("word", SPESE_TESTO))
             == "un elenco di 2 voci e una tabella di 3 righe con il totale")
    verifica("riassunto foglio", summary("excel", validate("excel", SPESE_FOGLIO))
             == "una tabella di 6 righe con il totale")
    new = json.loads(json.dumps(doc))
    new["blocchi"][0]["testo"] = "Roma, 15 ottobre 2026"
    verifica("cambiamento: un paragrafo", describe_change("word", doc, new) == "un paragrafo cambiato")
    lst = validate("pdf", {"titolo": "Compiti", "blocchi": [{"tipo": "elenco", "voci": ["a", "b"]}]})
    lst2 = json.loads(json.dumps(lst))
    lst2["blocchi"][0]["voci"].append("inglese")
    verifica("cambiamento: una voce in più", describe_change("pdf", lst, lst2) == "una voce in più")
    sh = validate("excel", SPESE_FOGLIO)
    sh2 = json.loads(json.dumps(sh))
    sh2["fogli"][0]["righe"][1][1] = 95
    verifica("cambiamento: una riga del foglio", describe_change("excel", sh, sh2) == "una riga cambiata")
    verifica("cambiamento: niente", describe_change("word", doc, doc) == "")


def prova_brain():
    """risposta_finale: la frase si dice e il turno finisce senza un'altra passata."""
    cfg, svc = servizio("brain")
    b = Brain.__new__(Brain)
    b.cfg, b.history = cfg, []
    b.tools = build_registry(documenti=svc.formati)
    b.tool_ctx = ToolContext(cfg=cfg, speakers=Speakers(), speaker_ctx=SpeakerCtx("Dario", "amministra"),
                             speaker=None, documenti=svc)

    class Finto:
        def __init__(self):
            self.chiamate = 0
            self.prompt = ""

        def stream(self, messages, tools):
            self.chiamate += 1
            self.prompt = messages[0]["content"]
            yield "calls", [{"id": "call_0", "name": "documento_crea",
                             "arguments": {"formato": "excel", "richiesta": "spese"}}]
    b.backend = Finto()
    detto = "".join(b.stream_reply("Fammi una tabella Excel delle spese", "amministra"))
    verifica("Brain: detta la frase del tool", detto.startswith("Ho preparato il foglio Excel"), detto)
    verifica("Brain: nessuna seconda passata", b.backend.chiamate == 1, str(b.backend.chiamate))
    verifica("Brain: storia chiusa con la risposta", [m["role"] for m in b.history]
             == ["user", "assistant", "tool", "assistant"])
    verifica("prompt: nomina i tool dei documenti", "documento_crea" in b.backend.prompt
             and "non usi internet e non comandi" in b.backend.prompt)
    b.record_announcement("Dario, è pronta la lettera.")
    verifica("annuncio attaccato all'ultima risposta", b.history[-1]["content"].endswith(
        "È pronta la lettera.") is False and "è pronta la lettera" in b.history[-1]["content"])
    verifica("prompt senza documenti non li nomina", "documento_crea" not in Config().prompt_for(False))
    svc.close()


def prova_formato_detto():
    """Il formato detto a voce vince solo se esplicito e non negato (rapporto del 01/10)."""
    from calliope.tools.documenti import _format
    for formato, frase, atteso in [
            ("word", "fammi un PDF con i compiti", "pdf"),             # il modello sbaglia, vince il detto
            ("excel", "scrivi una lettera su un foglio intestato", "word"),
            ("word", "scrivi una lettera su un foglio intestato", "word"),
            ("pdf", "una tabella delle spese, non in PDF", "pdf"),     # negato: decide il modello
            ("excel", "una tabella delle spese, non in PDF", "excel"),
            ("word", "fammi un foglio di calcolo con le spese", "excel"),
            ("", "una lettera di reclamo con le spese del condominio", "word"),
            ("", "una tabella delle spese di settembre", "excel"),
            ("excel", "fammi un PDF, non un Word", "pdf"),
            ("excel", "preparami la tabella delle spese", "excel")]:
        verifica(f"formato «{frase}» (modello: {formato or '-'}) → {atteso}",
                 _format(formato, frase) == atteso, _format(formato, frase))


def prova_config():
    text = example_yaml()
    verifica("calliope.yaml: sezione documenti", "documenti:" in text and "documenti_attesa_s: 4.0" in text
             and "- arial" in text)
    path = TMP / "c.yaml"
    path.write_text("documenti:\n  documenti_cartella: D:/Carte\n  documenti_font: [verdana]\n",
                    encoding="utf-8")
    import contextlib
    with contextlib.redirect_stdout(io.StringIO()):
        cfg = load_config(str(path))
    verifica("calliope.yaml: cartella e font letti", cfg.documenti_cartella == "D:/Carte"
             and cfg.documenti_font == ["verdana"])


prova_render()
prova_validazione()
prova_nomi()
prova_servizio()
prova_secondo_piano()
prova_rilettura()
prova_scrittore()
prova_frasi()
prova_brain()
prova_formato_detto()
prova_config()
print(f"\n{errori} errori" if errori else "\nTutto a posto.")
sys.exit(1 if errori else 0)
