"""
I modelli di documento: quelli già pronti (fattura, nota di credito, preventivo, DDT) e
quelli dell'utente nella cartella dei modelli (`ufficio_modelli`, predefinita
Documenti/Calliope/Modelli, fuori da git).

Un modello dell'utente è un file Word o PowerPoint con accanto un file di descrizione YAML
con lo stesso nome (Preventivo.docx + Preventivo.yaml):

    nome: preventivo sito             # come si dice a voce
    alias: [offerta sito]
    titolo: Preventivo per un sito web
    serie: preventivi                 # facoltativo: numerazione → {{ numero }}, {{ anno }}
    campi:
      cliente:   {tipo: anagrafica, descrizione: il cliente}
      oggetto:   {tipo: testo, descrizione: oggetto del preventivo}
      descrizione: {tipo: testo_lungo, obbligatorio: false,
                    descrizione: il lavoro, scritto per bene dallo scrittore}
      voci:      {tipo: righe, descrizione: le voci con quantità e prezzo}
      consegna:  {tipo: data, obbligatorio: false, predefinito: "da concordare"}

- Word (.docx, docxtpl): segnaposti Jinja2 dentro il documento, con stili, intestazione e
  carta intestata del file: {{ cliente.nome }}, {{ oggetto }}, righe di tabella ripetute
  con {%tr for v in voci %} … {%tr endfor %}, {{ totali.totale }}.
- PowerPoint (.pptx, python-pptx): gli stessi {{ … }} nelle caselle di testo, nelle tabelle
  e nelle note; una diapositiva con il commento {# ripeti voci come voce #} si ripete per
  ogni elemento; una riga di tabella con {# riga voci come v #} si ripete per ogni elemento.
- JSON (il formato dei modelli del 02/10, calliope/agenti/modelli.py): continua a valere,
  compilato con i documenti di sempre (Word, Excel o PDF).

Tipi dei campi: testo, testo_lungo (lo scrive lo scrittore partendo da quello che è stato
detto), numero, importo, data, booleano, elenco, tabella (con `colonne`), righe (voci con
quantità, prezzo e aliquota: i totali li calcola il programma), anagrafica (un nome della
rubrica: diventa il contatto intero).

Nei modelli Jinja gira in un ambiente **sandbox** con l'autoescape: un valore detto a voce
con «<» o «&» non rompe l'XML del file, e un modello non può chiamare codice.
"""

import copy
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

TIPI_CAMPO = ("testo", "testo_lungo", "numero", "importo", "data", "booleano", "elenco",
              "tabella", "righe", "anagrafica")
USCITE = {"word": "docx", "pdf": "pdf", "excel": "xlsx", "pptx": "pptx"}
_NUM = {"anyOf": [{"type": "number"}, {"type": "null"}]}
_STR = {"anyOf": [{"type": "string"}, {"type": "null"}]}


class ModelloNonValido(ValueError):
    pass


@dataclass
class Campo:
    nome: str
    tipo: str = "testo"
    obbligatorio: bool = True
    descrizione: str = ""
    colonne: list = field(default_factory=list)
    predefinito: object = None
    prezzi: bool = True              # per «righe»: con prezzo e aliquota (no per i DDT)
    chiedi: str = ""                 # come chiederlo a voce («il cliente»), se diverso

    def detto(self) -> str:
        return self.chiedi or self.descrizione or self.nome.replace("_", " ")

    def schema(self) -> dict:
        t = self.tipo
        if t in ("numero", "importo"):
            return _NUM
        if t == "booleano":
            return {"anyOf": [{"type": "boolean"}, {"type": "null"}]}
        if t == "elenco":
            return {"anyOf": [{"type": "array", "items": {"type": "string"}}, {"type": "null"}]}
        if t == "tabella":
            return {"anyOf": [{"type": "array", "items": {"type": "array", "items": {
                "anyOf": [{"type": "string"}, {"type": "number"}]}}}, {"type": "null"}]}
        if t == "righe":
            props = {"descrizione": {"type": "string"}, "quantita": _NUM}
            req = ["descrizione", "quantita"]
            if self.prezzi:
                props.update({"prezzo": _NUM, "aliquota": _NUM})
                req += ["prezzo", "aliquota"]
            else:
                props["unita"] = _STR
                req.append("unita")
            return {"anyOf": [{"type": "array", "items": {
                "type": "object", "properties": props, "required": req}}, {"type": "null"}]}
        return _STR


@dataclass
class Modello:
    nome: str
    titolo: str
    tipo: str                        # fattura | nota_di_credito | preventivo | ddt | docx | pptx | json
    uscita: str                      # word | pdf | excel | pptx
    campi: dict
    serie: str = ""                  # numerazione progressiva, se c'è
    alias: list = field(default_factory=list)
    descrizione: str = ""
    file: str = ""
    json_doc: dict | None = None     # i modelli JSON del 02/10
    avvisi: list = field(default_factory=list)
    fiscale: bool = False            # fattura e nota di credito (XML FatturaPA)

    def schema(self) -> dict:
        props = {n: c.schema() for n, c in self.campi.items()}
        return {"type": "object", "properties": props, "required": list(self.campi)}

    def descrivi_campi(self) -> str:
        return "; ".join(f"{n} ({c.tipo}{', obbligatorio' if c.obbligatorio else ''})"
                         + (f": {c.descrizione}" if c.descrizione else "")
                         for n, c in self.campi.items())

    def mancanti(self, dati: dict) -> list[str]:
        out = []
        for n, c in self.campi.items():
            v = (dati or {}).get(n)
            if c.obbligatorio and c.predefinito is None and (v is None or v == "" or v == []):
                out.append(n)
        return out

    def nomi_detti(self) -> list[str]:
        return [self.nome, *self.alias]


# ─────────────────────────── modelli pronti ───────────────────────────

def _righe(desc: str, prezzi: bool = True) -> Campo:
    return Campo("righe", "righe", True, desc, prezzi=prezzi,
                 chiedi="le voci con i prezzi" if prezzi else "i beni con le quantità")


def _cliente(desc: str = "il cliente (nome come in rubrica)", chiedi="il cliente") -> Campo:
    return Campo("cliente", "anagrafica", True, desc, chiedi=chiedi)


def pronti() -> dict[str, Modello]:
    fattura = Modello(
        "fattura", "fattura elettronica", "fattura", "pdf", {
            "cliente": _cliente(),
            "righe": _righe("le voci: descrizione, quantità (se detta), prezzo unitario come "
                            "detto e aliquota IVA solo se detta"),
            "iva_inclusa": Campo("iva_inclusa", "booleano", False,
                                 "true solo se i prezzi detti sono IVA inclusa"),
            "ritenuta": Campo("ritenuta", "booleano", False,
                              "true se chiede la ritenuta d'acconto, false se dice senza"),
            "data": Campo("data", "data", False, "data della fattura, se detta"),
            "scadenza_giorni": Campo("scadenza_giorni", "numero", False,
                                     "giorni per pagare, se detti"),
            "causale": Campo("causale", "testo", False, "causale o note, se dette"),
        }, serie="fatture", alias=["fatture"], fiscale=True,
        descrizione="PDF di cortesia e file XML FatturaPA da caricare allo SdI")
    nota = Modello(
        "nota di credito", "nota di credito", "nota_di_credito", "pdf", {
            "cliente": _cliente(),
            "fattura_collegata": Campo("fattura_collegata", "testo", True,
                                       "numero della fattura da correggere",
                                       chiedi="il numero della fattura da correggere"),
            "data_fattura_collegata": Campo("data_fattura_collegata", "data", False,
                                            "data della fattura da correggere"),
            "righe": _righe("le voci da stornare: descrizione, quantità, prezzo, aliquota"),
            "iva_inclusa": Campo("iva_inclusa", "booleano", False, "prezzi IVA inclusa"),
            "data": Campo("data", "data", False, "data, se detta"),
            "causale": Campo("causale", "testo", False, "motivo, se detto"),
        }, serie="note_credito", alias=["storno"], fiscale=True)
    prev = Modello(
        "preventivo", "preventivo", "preventivo", "pdf", {
            "cliente": _cliente(),
            "oggetto": Campo("oggetto", "testo", True, "oggetto del preventivo, in breve",
                             chiedi="l'oggetto"),
            "descrizione": Campo("descrizione", "testo_lungo", False,
                                 "cosa comprende il lavoro, come detto"),
            "righe": _righe("le voci: descrizione, quantità, prezzo unitario, aliquota"),
            "iva_inclusa": Campo("iva_inclusa", "booleano", False, "prezzi IVA inclusa"),
            "validita_giorni": Campo("validita_giorni", "numero", False,
                                     "per quanti giorni vale l'offerta"),
            "note": Campo("note", "testo", False, "altre condizioni dette"),
            "data": Campo("data", "data", False, "data, se detta"),
        }, serie="preventivi", alias=["offerta", "preventivi"])
    ddt = Modello(
        "ddt", "documento di trasporto", "ddt", "pdf", {
            "cliente": _cliente("il destinatario (nome come in rubrica)", "il destinatario"),
            "righe": _righe("i beni: descrizione, quantità e unità di misura", prezzi=False),
            "causale_trasporto": Campo("causale_trasporto", "testo", False,
                                       "vendita, conto visione, reso, riparazione…"),
            "colli": Campo("colli", "numero", False, "numero di colli"),
            "peso": Campo("peso", "testo", False, "peso, come detto"),
            "aspetto": Campo("aspetto", "testo", False, "aspetto dei beni (scatole, pallet…)"),
            "trasporto_a_cura": Campo("trasporto_a_cura", "testo", False,
                                      "mittente, destinatario o vettore"),
            "vettore": Campo("vettore", "testo", False, "nome del corriere"),
            "luogo_destinazione": Campo("luogo_destinazione", "testo", False,
                                        "se diverso dall'indirizzo del destinatario"),
            "data": Campo("data", "data", False, "data, se detta"),
        }, serie="ddt", alias=["documento di trasporto", "bolla", "bolla di consegna"])
    return {m.nome: m for m in (fattura, nota, prev, ddt)}


# ─────────────────────────── modelli dell'utente ───────────────────────────

def _campi_da(raw, nome_modello: str) -> dict:
    campi = {}
    for n, c in (raw or {}).items():
        if not re.fullmatch(r"[a-zA-Z_]\w*", str(n)):
            raise ModelloNonValido(f"{nome_modello}: nome di campo non valido {n!r}")
        c = c if isinstance(c, dict) else {"descrizione": str(c)}
        tipo = str(c.get("tipo") or "testo")
        if tipo not in TIPI_CAMPO:
            raise ModelloNonValido(f"{nome_modello}: tipo {tipo!r} del campo {n} sconosciuto "
                                   f"(ammessi: {', '.join(TIPI_CAMPO)})")
        campi[str(n)] = Campo(str(n), tipo, bool(c.get("obbligatorio", True)),
                              str(c.get("descrizione") or ""), list(c.get("colonne") or []),
                              c.get("predefinito"), bool(c.get("prezzi", True)),
                              str(c.get("chiedi") or ""))
    if not campi:
        raise ModelloNonValido(f"{nome_modello}: nessun campo")
    return campi


def _segnaposti_docx(path: Path) -> set[str]:
    """I nomi usati nei segnaposti del file (per avvisare di campi senza segnaposto)."""
    import zipfile
    testo = ""
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith(("word/", "ppt/")) and n.endswith(".xml"):
                testo += re.sub(r"<[^>]+>", "", z.read(n).decode("utf-8", "replace"))
    return set(re.findall(r"\{\{\s*([a-zA-Z_]\w*)", testo)) | set(
        re.findall(r"\b(?:in|ripeti|riga)\s+([a-zA-Z_]\w*)", testo))


def leggi_modello(descrizione: Path) -> Modello:
    """Un modello dell'utente dal suo file di descrizione YAML (accanto al .docx o .pptx)."""
    import yaml
    descrizione = Path(descrizione)
    try:
        data = yaml.safe_load(descrizione.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as e:
        raise ModelloNonValido(f"{descrizione.name} non si legge: {type(e).__name__}") from None
    if not isinstance(data, dict):
        raise ModelloNonValido(f"{descrizione.name}: serve un oggetto YAML")
    nome = str(data.get("nome") or descrizione.stem).strip().lower()
    file = None
    for ext in ("docx", "pptx"):
        cand = descrizione.with_suffix("." + ext)
        if cand.is_file():
            file = cand
    if data.get("file"):
        file = descrizione.parent / str(data["file"])
    if file is None or not file.is_file():
        raise ModelloNonValido(f"{nome}: manca il file Word o PowerPoint accanto a "
                               f"{descrizione.name}")
    tipo = file.suffix.lower().lstrip(".")
    if tipo not in ("docx", "pptx"):
        raise ModelloNonValido(f"{nome}: il file deve essere .docx o .pptx")
    campi = _campi_da(data.get("campi"), nome)
    serie = str(data.get("serie") or "").strip().lower()
    if serie and not re.fullmatch(r"[a-z_]{2,30}", serie):
        raise ModelloNonValido(f"{nome}: serie {serie!r} non valida (lettere e _)")
    m = Modello(nome, str(data.get("titolo") or nome), tipo,
                "word" if tipo == "docx" else "pptx", campi, serie,
                [str(a).lower() for a in data.get("alias") or []],
                str(data.get("descrizione") or ""), str(file))
    try:
        usati = _segnaposti_docx(file)
        m.avvisi = [f"campo {n} non usato nel file" for n in campi if n not in usati]
    except Exception as e:  # noqa: BLE001 — un file rovinato si scopre alla compilazione
        m.avvisi = [f"file non leggibile: {type(e).__name__}"]
    return m


def dal_json(path: Path) -> Modello:
    """Un modello JSON del 02/10 (calliope/agenti/modelli.py) nel catalogo nuovo."""
    from ..agenti.modelli import ModelloNonValido as Vecchio
    from ..agenti.modelli import leggi_modello as leggi_vecchio
    try:
        v = leggi_vecchio(path)
    except Vecchio as e:
        raise ModelloNonValido(str(e)) from None
    tipo_map = {"testo": "testo", "numero": "numero", "data": "data", "elenco": "elenco",
                "tabella": "tabella"}
    campi = {n: Campo(n, tipo_map.get(c.tipo, "testo"), c.obbligatorio, c.descrizione)
             for n, c in v.campi.items()}
    return Modello(v.nome.lower(), v.titolo, "json", v.formato, campi, file=str(path),
                   json_doc=v.documento, descrizione=v.descrizione, avvisi=list(v.avvisi))


def carica(cartella) -> tuple[dict[str, Modello], list[str]]:
    """(modelli pronti + quelli della cartella, errori dei modelli rovinati). Un modello
    dell'utente con lo stesso nome di uno pronto lo sostituisce (la propria fattura no:
    fattura e nota di credito restano quelle con l'XML)."""
    out = pronti()
    errori = []
    p = Path(cartella) if cartella else None
    if p is None or not p.is_dir():
        return out, errori
    for f in sorted(p.iterdir()):
        try:
            if f.suffix.lower() in (".yaml", ".yml"):
                m = leggi_modello(f)
            elif f.suffix.lower() == ".json":
                m = dal_json(f)
            else:
                continue
        except ModelloNonValido as e:
            errori.append(str(e))
            continue
        if m.nome in out and out[m.nome].fiscale:
            errori.append(f"{m.nome}: nome riservato alla fattura elettronica, rinominalo")
            continue
        out[m.nome] = m
    return out, errori


def risolvi(catalogo: dict[str, Modello], detto: str) -> Modello | None:
    """Il modello dal nome detto (o dall'alias), anche con un articolo o un plurale."""
    import difflib
    t = re.sub(r"^(il|lo|la|un|una|uno|l')\s*", "", str(detto or "").strip().lower())
    for m in catalogo.values():
        if t in m.nomi_detti():
            return m
    nomi = {n: m for m in catalogo.values() for n in m.nomi_detti()}
    vicini = difflib.get_close_matches(t, list(nomi), n=1, cutoff=0.8)
    return nomi[vicini[0]] if vicini else None


# ─────────────────────────── compilazione: Word ───────────────────────────

def _jinja():
    from jinja2.sandbox import SandboxedEnvironment
    return SandboxedEnvironment(autoescape=False)


def compila_docx(path: str, contesto: dict) -> bytes:
    """Il .docx dell'utente con i dati (docxtpl). autoescape: «&» e «<» nei valori."""
    from docxtpl import DocxTemplate
    tpl = DocxTemplate(path)
    tpl.render(contesto, jinja_env=_jinja(), autoescape=True)
    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


# ─────────────────────────── compilazione: PowerPoint ───────────────────────────
_RIPETI = re.compile(r"\{#\s*ripeti\s+([a-zA-Z_][\w.]*)(?:\s+come\s+([a-zA-Z_]\w*))?\s*#\}")
_RIGA = re.compile(r"\{#\s*riga\s+([a-zA-Z_][\w.]*)(?:\s+come\s+([a-zA-Z_]\w*))?\s*#\}")


def _testi(slide):
    """I text frame della diapositiva: caselle, celle delle tabelle, gruppi e note."""
    def da(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                       # gruppo
                yield from da(sh.shapes)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                yield sh.text_frame
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame
    yield from da(slide.shapes)
    if slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame


def _sostituisci(tf, env, ctx):
    """Ogni paragrafo con un segnaposto si riscrive nel suo primo pezzo (run): la
    formattazione del primo pezzo vale per tutto il paragrafo."""
    for p in tf.paragraphs:
        runs = p.runs
        testo = "".join(r.text for r in runs)
        if "{{" not in testo and "{#" not in testo and "{%" not in testo:
            continue
        nuovo = env.from_string(testo).render(**ctx)
        if runs:
            runs[0].text = nuovo
            for r in runs[1:]:
                r.text = ""


def _valore(ctx: dict, percorso: str):
    v = ctx
    for parte in percorso.split("."):
        v = v.get(parte) if isinstance(v, dict) else getattr(v, parte, None)
    return v or []


def _duplica(prs, slide):
    """Una copia della diapositiva subito dopo l'originale (forme, immagini, note escluse)."""
    nuova = prs.slides.add_slide(slide.slide_layout)
    for sh in list(nuova.shapes):
        sh._element.getparent().remove(sh._element)
    ids = {}
    for rid, rel in slide.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            ids[rid] = nuova.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            ids[rid] = nuova.part.relate_to(rel.target_part, rel.reltype)
    for sh in slide.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr, val in list(node.attrib.items()):
                if attr.endswith("}embed") or attr.endswith("}link") or attr.endswith("}id"):
                    if val in ids:
                        node.set(attr, ids[val])
        nuova.shapes._spTree.append(el)
    lista = prs.slides._sldIdLst
    elems = list(lista)
    nuovo_el = elems[-1]
    lista.remove(nuovo_el)
    pos = elems.index(next(e for e in elems if prs.slides.part.related_part(
        e.rId) is slide.part))
    lista.insert(pos + 1, nuovo_el)
    return nuova


def _togli(prs, slide):
    lista = prs.slides._sldIdLst
    for e in list(lista):
        if prs.slides.part.related_part(e.rId) is slide.part:
            prs.part.drop_rel(e.rId)
            lista.remove(e)


def _ripeti_righe(slide, env, ctx):
    for sh in slide.shapes:
        if not getattr(sh, "has_table", False) or not sh.has_table:
            continue
        tbl = sh.table._tbl
        for tr in list(tbl.tr_lst):
            testo = "".join(t.text or "" for t in tr.iter() if t.tag.endswith("}t"))
            m = _RIGA.search(testo)
            if not m:
                continue
            elementi = _valore(ctx, m.group(1))
            nome = m.group(2) or "elemento"
            pos = list(tbl).index(tr)
            for k, el in enumerate(elementi):
                riga = copy.deepcopy(tr)
                for t in riga.iter():
                    if t.tag.endswith("}t") and t.text:
                        t.text = _RIGA.sub("", t.text)
                # Le celle si compilano come testo semplice con l'elemento corrente
                for t in [x for x in riga.iter() if x.tag.endswith("}t") and x.text and
                          ("{{" in x.text or "{%" in x.text)]:
                    t.text = env.from_string(t.text).render(**dict(ctx, **{nome: el}))
                tbl.insert(pos + 1 + k, riga)
            tbl.remove(tr)


def compila_pptx(path: str, contesto: dict) -> bytes:
    from pptx import Presentation
    prs = Presentation(path)
    env = _jinja()
    for slide in list(prs.slides):
        testo = " ".join(p.text for tf in _testi(slide) for p in tf.paragraphs)
        m = _RIPETI.search(testo)
        if not m:
            _ripeti_righe(slide, env, contesto)
            for tf in _testi(slide):
                _sostituisci(tf, env, contesto)
            continue
        elementi = list(_valore(contesto, m.group(1)))
        nome = m.group(2) or "elemento"
        if not elementi:
            _togli(prs, slide)
            continue
        copie = [slide]
        for _ in elementi[1:]:
            copie.append(_duplica(prs, copie[-1]))
        for s, el in zip(copie, elementi):
            ctx = dict(contesto, **{nome: el})
            _ripeti_righe(s, env, ctx)
            for tf in _testi(s):
                _sostituisci(tf, env, ctx)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def compila_json(m: Modello, dati: dict) -> dict:
    """I modelli JSON del 02/10: il documento a blocchi con i dati (agenti/modelli.py)."""
    from ..agenti.modelli import Campo as CampoV
    from ..agenti.modelli import Modello as ModelloV
    v = ModelloV(m.nome, m.titolo, m.uscita,
                 {n: CampoV(n, c.descrizione, c.tipo if c.tipo in ("testo", "numero", "data",
                                                                    "elenco", "tabella")
                            else "testo", c.obbligatorio) for n, c in m.campi.items()},
                 m.json_doc or {})
    piatti = {k: (v_["nome_completo"] if isinstance(v_, dict) and "nome_completo" in v_ else v_)
              for k, v_ in dati.items()}
    return v.compila(piatti)


def librerie() -> dict[str, bool]:
    """Quali compilazioni si possono fare qui (docxtpl per Word, python-pptx per PowerPoint)."""
    from importlib.util import find_spec
    return {"docx": find_spec("docxtpl") is not None, "pptx": find_spec("pptx") is not None}
