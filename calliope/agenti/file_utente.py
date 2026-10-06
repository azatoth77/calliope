"""
I file della persona dati all'agente (03/10/2026): «correggi lo script backup.py»,
«riassumimi il PDF del contratto», «aggiungi una colonna al foglio spese.xlsx».

Il flusso (calliope/tools/agenti.py e servizio.py):
1. `delega_lavoro(file=…)` cerca il file sul PC con l'esecutore (lo stesso di pc_cerca_file,
   con gli stessi permessi: proprietari del PC o chi amministra, oppure un documento appena
   scritto per chi parla) e propone «Mando una copia di … all'agente … Procedo?»: il file
   lascia il portatile, quindi la conferma è sempre esplicita (azione in sospeso);
2. al «sì» la copia si prende in secondo piano (`PCExecutor.copia_file`: in locale si legge
   il file, con il satellite arriva a pezzi con lo SHA-256), con un tetto alla dimensione
   (`agenti_file_max_mb`) e solo le estensioni di `ESTENSIONI` (testo, codice, Word, Excel,
   PDF); il satellite ricontrolla tutto per conto suo;
3. il codice lavora sulla copia nella sandbox; per gli altri lavori il testo del file
   (`testo_del_file`) entra nella richiesta all'agente;
4. il risultato torna come file **nuovo** («backup (corretto).py», «spese (modificato).xlsx»:
   `nome_risultato`), con `RemoteDelivery` nella cartella Calliope dei Documenti del
   portatile; l'originale non si tocca mai.
"""

import io
import re
from pathlib import PurePath

# Testo e codice: l'agente li legge e li riscrive nella sandbox
ESTENSIONI_TESTO = ("py", "txt", "md", "csv", "tsv", "json", "html", "htm", "css", "js",
                    "ps1", "psm1", "sql", "toml", "yaml", "yml", "ini", "cfg", "xml", "log")
# Documenti: il testo per i lavori di documento, il file intero (openpyxl, python-docx) per
# il codice
ESTENSIONI_DOCUMENTO = ("docx", "xlsx", "pdf")
ESTENSIONI = ESTENSIONI_TESTO + ESTENSIONI_DOCUMENTO

# Come si dice il tipo di file a voce: (articolo e nome, con «di» davanti)
_DETTO = {"py": ("lo script Python", "dello script Python"),
          "ps1": ("lo script PowerShell", "dello script PowerShell"),
          "pdf": ("il PDF", "del PDF"), "docx": ("il documento Word", "del documento Word"),
          "xlsx": ("il foglio Excel", "del foglio Excel"), "csv": ("il file CSV", "del file CSV"),
          "txt": ("il file di testo", "del file di testo"),
          "md": ("il file di testo", "del file di testo"),
          "json": ("il file JSON", "del file JSON"),
          "html": ("la pagina web", "della pagina web"),
          "htm": ("la pagina web", "della pagina web")}


def detto(nome: str, estensione: str, di: bool = False) -> str:
    """«il PDF «Contratto affitto»», «dello script Python «backup»» (con `di`): per la voce."""
    forme = _DETTO.get(estensione.lower(), ("il file", "del file"))
    return f"{forme[1] if di else forme[0]} «{nome}»"


def estensione(nome: str) -> str:
    return PurePath(str(nome or "")).suffix.lower().lstrip(".")


def nome_sicuro(nome: str) -> str:
    """Il nome di un file della persona come nome nella sandbox: solo la parte finale,
    caratteri semplici (la sandbox rifiuta gli altri), estensione conservata."""
    p = PurePath(str(nome or "").replace("\\", "/"))
    stem, ext = p.stem, p.suffix.lower()
    stem = re.sub(r"[^\w\-. ()+,=@]", "_", stem).strip(" .") or "file"
    if stem.startswith("."):
        stem = "_" + stem.lstrip(".")
    return f"{stem[:80]}{ext}"


def nome_risultato(nome: str, tipo: str) -> str:
    """«backup.py» → «backup (corretto).py» (codice), «spese.xlsx» → «spese
    (modificato).xlsx» (gli altri lavori): il risultato non si confonde con l'originale."""
    p = PurePath(nome)
    parola = "corretto" if tipo == "codice" and p.suffix.lower().lstrip(".") in \
        ESTENSIONI_TESTO else "modificato"
    return f"{p.stem} ({parola}){p.suffix}"


class FileNonLeggibile(Exception):
    """Il testo del file non si estrae (libreria mancante, file rovinato)."""


def testo_del_file(nome: str, dati: bytes, max_caratteri: int = 40_000) -> str:
    """Il testo di un file per l'agente: testo e codice decodificati, Word paragrafo per
    paragrafo (con le tabelle), Excel foglio per foglio come righe separate da «;», PDF
    pagina per pagina (pypdf, puro Python). Tagliato a `max_caratteri`."""
    ext = estensione(nome)
    if ext in ESTENSIONI_TESTO:
        for enc in ("utf-8-sig", "cp1252"):
            try:
                testo = dati.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            testo = dati.decode("utf-8", errors="replace")
    elif ext == "docx":
        try:
            import docx
        except ImportError as e:
            raise FileNonLeggibile("manca python-docx per leggere i file Word") from e
        try:
            d = docx.Document(io.BytesIO(dati))
        except Exception as e:  # noqa: BLE001 — file rovinato o non Word
            raise FileNonLeggibile("il file Word non si apre") from e
        righe = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                righe.append(" ; ".join(c.text.strip() for c in r.cells))
        testo = "\n".join(righe)
    elif ext == "xlsx":
        try:
            import openpyxl
        except ImportError as e:
            raise FileNonLeggibile("manca openpyxl per leggere i file Excel") from e
        try:
            wb = openpyxl.load_workbook(io.BytesIO(dati), read_only=True, data_only=False)
        except Exception as e:  # noqa: BLE001
            raise FileNonLeggibile("il file Excel non si apre") from e
        righe = []
        for ws in wb.worksheets:
            righe.append(f"[foglio {ws.title}]")
            for r in ws.iter_rows(values_only=True):
                if any(v is not None for v in r):
                    righe.append(" ; ".join("" if v is None else str(v) for v in r))
                if sum(len(x) for x in righe) > max_caratteri:
                    break
        wb.close()
        testo = "\n".join(righe)
    elif ext == "pdf":
        try:
            import pypdf
        except ImportError as e:
            raise FileNonLeggibile("manca pypdf per leggere i PDF (pip install pypdf)") from e
        try:
            r = pypdf.PdfReader(io.BytesIO(dati))
            pagine = []
            for i, pg in enumerate(r.pages, 1):
                pagine.append(f"[pagina {i}]\n{pg.extract_text() or ''}")
                if sum(len(x) for x in pagine) > max_caratteri:
                    break
        except Exception as e:  # noqa: BLE001
            raise FileNonLeggibile("il PDF non si legge") from e
        testo = "\n".join(pagine)
        if not re.search(r"\w{3}", re.sub(r"\[pagina \d+\]", "", testo)):
            raise FileNonLeggibile("il PDF non ha testo (è una scansione?)")
    else:
        raise FileNonLeggibile(f"non so leggere i file .{ext}")
    testo = testo.replace("\r\n", "\n").strip()
    if len(testo) > max_caratteri:
        testo = testo[:max_caratteri] + f"\n[… tagliato: il file è più lungo]"
    return testo
