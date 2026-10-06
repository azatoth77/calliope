"""
I tre renderer: dal JSON validato (formato.py) ai byte del file.

  docx → python-docx     xlsx → openpyxl     pdf → fpdf2 (puro Python)

Restituiscono byte, non scrivono file: la consegna (consegna.py) è un passo a parte, così
domani il file può andare a un PC remoto senza toccare questa parte.

Il PDF ha bisogno di un font TrueType per accenti, «€» e virgolette tipografiche: i font
base del PDF (Helvetica) conoscono solo il latin-1. Si cerca tra i font di sistema
(`Config.documenti_font`: nomi o percorsi, il primo che c'è); senza, si ripiega su
Helvetica e si traslitterano i caratteri che mancano («€» → «EUR»).
"""

import io
import os
import sys
from pathlib import Path

from .formato import (classify_cell, col_letter, format_number, is_letter, numeric_columns,
                      parse_number, table_total)

# ─────────────────────────── font per il PDF ───────────────────────────


def font_dirs() -> list[Path]:
    """Dove stanno i font di sistema (Windows, poi Linux e macOS per completezza)."""
    dirs = []
    if sys.platform == "win32":
        windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot") or r"C:\Windows"
        dirs.append(Path(windir) / "Fonts")
        local = os.environ.get("LOCALAPPDATA")
        if local:
            dirs.append(Path(local) / "Microsoft" / "Windows" / "Fonts")
    else:
        dirs += [Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
                 Path.home() / ".fonts", Path("/Library/Fonts"), Path("/System/Library/Fonts")]
    return [d for d in dirs if d.is_dir()]


def _find_file(name: str, dirs: list[Path]) -> Path | None:
    for d in dirs:
        for ext in ("", ".ttf", ".TTF"):
            p = d / f"{name}{ext}"
            if p.is_file():
                return p
    for d in dirs:                                   # Linux: i font stanno in sottocartelle
        if sys.platform != "win32":
            for p in d.rglob(f"{name}.ttf"):
                return p
    return None


def find_font(candidates) -> tuple[str, str | None] | None:
    """(regolare, grassetto o None) per il primo candidato che c'è. Un candidato è un
    percorso completo o un nome di file senza estensione («arial», «DejaVuSans»); il
    grassetto si cerca accanto con i nomi soliti (arialbd, segoeuib, DejaVuSans-Bold)."""
    dirs = font_dirs()
    for cand in candidates or []:
        cand = str(cand).strip()
        if not cand:
            continue
        path = Path(cand)
        regular = path if path.is_file() else _find_file(cand, dirs)
        if regular is None:
            continue
        stem, folder = regular.stem, [regular.parent] + dirs
        bold = next((b for b in (_find_file(stem + suffix, folder)
                                 for suffix in ("bd", "b", "-Bold", " Bold", "_Bold"))
                     if b is not None), None)
        return str(regular), (str(bold) if bold else None)
    return None


# Senza TrueType: Helvetica sa solo il latin-1 (le lettere accentate sì, «€» e le
# virgolette tipografiche no)
_LATIN1 = {"€": "EUR", "“": '"', "”": '"', "„": '"', "‘": "'", "’": "'", "–": "-", "—": "-",
           "…": "...", "•": "-", "\u00a0": " ", "\u202f": " "}


def _latin1(text: str) -> str:
    for a, b in _LATIN1.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


# ─────────────────────────── Word ───────────────────────────

def _table_rows(block: dict) -> tuple[list, list, list | None]:
    cols, rows = block["colonne"], block["righe"]
    return cols, rows, (table_total(cols, rows) if block.get("totale") else None)


def _blocks(doc: dict) -> list[dict]:
    """I blocchi da stampare, con il titolo del documento in cima se il contenuto non ne
    ha già uno. Il 01/10 «metti in alto il titolo Compiti Matteo» cambiava solo
    doc["titolo"], che finiva nelle proprietà del file: il PDF riscritto sembrava uguale.
    Le lettere no: lì un'intestazione grande non ci va (l'oggetto è nel testo)."""
    blocks = doc["blocchi"]
    if any(b["tipo"] == "titolo" for b in blocks) or is_letter(doc["titolo"]):
        return blocks
    return [{"tipo": "titolo", "testo": doc["titolo"]}] + blocks


def render_docx(doc: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    d.core_properties.title = doc["titolo"]
    d.core_properties.author = "Calliope"
    d.core_properties.language = "it-IT"
    first_heading = True
    for b in _blocks(doc):
        kind = b["tipo"]
        if kind == "titolo":
            d.add_heading(b["testo"], level=1 if first_heading else 2)
            first_heading = False
        elif kind == "paragrafo":
            p = d.add_paragraph()
            lines = b["testo"].split("\n")
            run = p.add_run(lines[0])
            for line in lines[1:]:                   # a capo dentro il paragrafo (w:br)
                run.add_break()
                run.add_text(line)
            if b.get("allinea") == "destra":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif kind == "elenco":
            style = "List Number" if b.get("numerato") else "List Bullet"
            for v in b["voci"]:
                d.add_paragraph(v, style=style)
        elif kind == "tabella":
            cols, rows, total = _table_rows(b)
            numeric = numeric_columns(cols, rows)
            all_rows = [cols] + rows + ([total] if total else [])
            table = d.add_table(rows=len(all_rows), cols=len(cols))
            table.style = "Table Grid"
            for r, row in enumerate(all_rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = str(value)
                    para = cell.paragraphs[0]
                    if r == 0 or (total and r == len(all_rows) - 1):
                        for run in para.runs:
                            run.bold = True
                    if r > 0 and c in numeric:
                        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            d.add_paragraph()
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ─────────────────────────── PDF ───────────────────────────

def render_pdf(doc: dict, font_candidates=None) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=20)
    pdf.set_title(doc["titolo"])
    pdf.set_author("Calliope")
    pdf.set_lang("it")
    font = find_font(font_candidates)
    if font:
        pdf.add_font("Testo", "", font[0])
        if font[1]:
            pdf.add_font("Testo", "B", font[1])
        family, bold, clean = "Testo", ("B" if font[1] else ""), (lambda s: s)
    else:
        family, bold, clean = "Helvetica", "B", _latin1
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    first_heading = True
    for b in _blocks(doc):
        kind = b["tipo"]
        if kind == "titolo":
            pdf.set_font(family, bold, 16 if first_heading else 13)
            pdf.multi_cell(width, 8, clean(b["testo"]), align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            first_heading = False
        elif kind == "paragrafo":
            pdf.set_font(family, "", 11)
            align = "R" if b.get("allinea") == "destra" else "L"
            pdf.multi_cell(width, 6, clean(b["testo"]), align=align, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)
        elif kind == "elenco":
            pdf.set_font(family, "", 11)
            for i, v in enumerate(b["voci"], 1):
                mark = f"{i}." if b.get("numerato") else clean("•")
                pdf.set_x(pdf.l_margin + 2)
                pdf.cell(7, 6, mark)
                pdf.multi_cell(width - 9, 6, clean(v), align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)
        elif kind == "tabella":
            cols, rows, total = _table_rows(b)
            numeric = numeric_columns(cols, rows)
            pdf.set_font(family, "", 10)
            head = FontFace(emphasis=bold or None, fill_color=(230, 230, 230))
            align = tuple("RIGHT" if c in numeric else "LEFT" for c in range(len(cols)))
            with pdf.table(headings_style=head, text_align=align, line_height=6,
                           width=width) as table:
                for row in [cols] + rows:
                    r = table.row()
                    for value in row:
                        r.cell(clean(str(value)))
                if total:
                    r = table.row()
                    for value in total:
                        r.cell(clean(str(value)), style=FontFace(emphasis=bold or None))
            pdf.ln(4)
    return bytes(pdf.output())


# ─────────────────────────── Excel ───────────────────────────

_EURO_FORMAT = '#,##0.00 "€"'


def render_xlsx(doc: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = doc["titolo"]
    wb.properties.creator = "Calliope"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E6E6E6")
    for s in doc["fogli"]:
        ws = wb.create_sheet(s["nome"])
        cols, rows = s["colonne"], s["righe"]
        ncols, nrows = len(cols), len(rows)
        max_row = nrows + 1 + (1 if s.get("totale") else 0)
        euro_cols = {c for c, euro in numeric_columns(cols, rows).items() if euro}
        widths = [len(str(c)) for c in cols]
        for c, name in enumerate(cols, 1):
            cell = ws.cell(row=1, column=c, value=str(name))
            cell.font, cell.fill = bold, fill
        for r, row in enumerate(rows, 2):
            for c, value in enumerate(row, 1):
                kind = classify_cell(value, ncols, max_row)
                cell = ws.cell(row=r, column=c)
                if kind[0] == "formula":
                    cell.value = kind[1]
                    if c - 1 in euro_cols:
                        cell.number_format = _EURO_FORMAT
                elif kind[0] == "numero":
                    cell.value = kind[1]
                    if kind[2] or c - 1 in euro_cols:
                        cell.number_format = _EURO_FORMAT
                elif kind[0] == "testo":
                    cell.value = kind[1]
                    # Testo puro, mai formula: «=…», «+39…», «@…» non devono diventare
                    # formule quando qualcuno apre o modifica la cella (formula injection)
                    cell.data_type = "s"
                    if kind[2]:
                        cell.quotePrefix = True
                widths[c - 1] = max(widths[c - 1], len(str(value)))
        if s.get("totale"):
            r = nrows + 2
            label = ws.cell(row=r, column=1, value="Totale")
            label.font = bold
            for c in numeric_columns(cols, rows):
                letter = col_letter(c + 1)
                cell = ws.cell(row=r, column=c + 1,
                               value=f"=SUM({letter}2:{letter}{nrows + 1})" if nrows else 0)
                cell.font = bold
                if c in euro_cols:
                    cell.number_format = _EURO_FORMAT
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[col_letter(c)].width = min(50, max(10, w + 2))
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    cell.alignment = Alignment(wrap_text=len(cell.value) > 50, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sheet_totals(doc: dict) -> list[tuple[str, str]]:
    """I totali di ogni foglio, calcolati qui (per le prove e per chi volesse dirli)."""
    out = []
    for s in doc.get("fogli", []):
        if not s.get("totale"):
            continue
        for c, euro in numeric_columns(s["colonne"], s["righe"]).items():
            total = sum(parse_number(r[c])[0] for r in s["righe"]
                        if parse_number(r[c]) is not None)
            out.append((s["colonne"][c], format_number(round(total, 2), euro)))
    return out


RENDERERS = {"word": render_docx, "excel": render_xlsx, "pdf": render_pdf}


def available_formats() -> tuple[tuple[str, ...], dict[str, str]]:
    """(formati disponibili, {formato: libreria mancante}). Senza una libreria quel
    formato semplicemente non c'è."""
    ok, missing = [], {}
    for formato, module, package in (("word", "docx", "python-docx"),
                                     ("excel", "openpyxl", "openpyxl"),
                                     ("pdf", "fpdf", "fpdf2")):
        try:
            __import__(module)
            ok.append(formato)
        except ImportError:
            missing[formato] = package
    return tuple(ok), missing


def render(formato: str, doc: dict, font_candidates=None) -> bytes:
    if formato == "pdf":
        return render_pdf(doc, font_candidates)
    return RENDERERS[formato](doc)


def plain_text(formato: str, data: bytes) -> str:
    """Il testo di un file generato, riletto con le stesse librerie (per le prove)."""
    if formato == "word":
        from docx import Document
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if formato == "excel":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        return "\n".join(" | ".join("" if v is None else str(v) for v in row)
                         for ws in wb for row in ws.iter_rows(values_only=True))
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
