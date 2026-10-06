"""
Il testo dei documenti di casa (03/10/2026): PDF con testo, PDF scansionati, foto e
immagini, Word, testo semplice.

- **PDF**: pypdfium2 (PDFium, wheel per win_amd64, win_arm64 e manylinux aarch64) legge il
  testo di ogni pagina; una pagina quasi senza testo (scansione) si disegna come immagine e
  va all'OCR;
- **immagini** (JPG, PNG, WebP, TIFF anche a più pagine): Pillow, con l'orientamento EXIF
  delle foto del telefono, ridotte a `lato_max` pixel;
- **Word** (.docx): python-docx, paragrafi e tabelle;
- **OCR**: un modello multimodale già servito (qwen3.6 su vLLM sulla DGX, o un Ollama con un
  modello visivo). Scelto misurando contro RapidOCR in ONNX (docs/ricerche/
  2026-10-03-documenti-grafo.md, §3): sulle foto peggiori il modello visivo legge 29 valori
  chiave su 30, RapidOCR 17, e non servono librerie native nuove. Costa ~4–5 s a pagina,
  in secondo piano.

`VERSIONE` cambia quando cambia il modo di leggere: allora si rifà anche l'OCR (altrimenti,
a estrattore nuovo, si riusa il testo salvato).
"""

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path

VERSIONE = "1"

ESTENSIONI = {".pdf": "pdf", ".jpg": "immagine", ".jpeg": "immagine", ".png": "immagine",
              ".webp": "immagine", ".tif": "immagine", ".tiff": "immagine",
              ".bmp": "immagine", ".docx": "docx", ".txt": "testo"}

PROMPT_OCR = ("Trascrivi fedelmente tutto il testo di questa immagine di un documento, riga "
              "per riga, nell'ordine di lettura; in una tabella tieni ogni voce sulla stessa "
              "riga del suo importo. Non correggere, non riassumere, non aggiungere commenti "
              "né markdown: solo il testo. Se non c'è testo, rispondi con una riga vuota.")


class ErroreTesto(Exception):
    """Il file non si legge (libreria mancante, file rovinato, OCR spento): il messaggio
    dice perché, in italiano."""


@dataclass
class Estratto:
    testo: str
    metodo: str                     # «pdf», «ocr», «pdf+ocr», «docx», «testo»
    pagine: int = 1
    pagine_ocr: int = 0
    secondi_ocr: float = 0.0
    avvisi: list = field(default_factory=list)


class OcrVisivo:
    """OCR con un modello multimodale: API compatibile OpenAI (vLLM, llama.cpp) o nativa di
    Ollama («images»). Chiamato solo dal thread dell'archivio."""

    def __init__(self, url: str, modello: str, motore: str = "openai", timeout: float = 180.0,
                 chiave: str | None = None, prima_di_chiamare=None, voce: dict | None = None):
        import httpx
        u = url.rstrip("/")
        if u.endswith("/v1"):
            u = u[:-3]
        self.motore = motore
        self.modello = modello
        self.prima = prima_di_chiamare
        # Stesso Ollama e stesso modello della voce: num_ctx e keep_alive della voce
        # (agenti.impostazioni.opzioni_voce), altrimenti Ollama ricarica il modello
        self.voce = dict(voce or {})
        headers = {"Authorization": f"Bearer {chiave}"} if chiave else {}
        self.http = httpx.Client(base_url=u, headers=headers,
                                 timeout=httpx.Timeout(timeout, connect=10.0))

    def leggi(self, img, fonte: str = "") -> str:
        """Il testo dell'immagine. `fonte` («file.pdf, pagina 2») serve solo ai messaggi."""
        if self.prima is not None:
            self.prima()
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=88)
        b64 = base64.b64encode(buf.getvalue()).decode()
        if self.motore == "ollama":
            body = {"model": self.modello, "stream": False, "think": False,
                    "options": {"temperature": 0, "num_predict": 3000,
                                **self.voce.get("options", {})},
                    "messages": [{"role": "user", "content": PROMPT_OCR, "images": [b64]}]}
            if "keep_alive" in self.voce:
                body["keep_alive"] = self.voce["keep_alive"]
            r = self.http.post("/api/chat", json=body)
            r.raise_for_status()
            return (r.json().get("message") or {}).get("content") or ""
        body = {"model": self.modello, "temperature": 0, "max_tokens": 3000,
                "chat_template_kwargs": {"enable_thinking": False},
                "messages": [{"role": "user", "content": [
                    {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + b64}},
                    {"type": "text", "text": PROMPT_OCR}]}]}
        r = self.http.post("/v1/chat/completions", json=body)
        r.raise_for_status()
        return r.json()["choices"][0]["message"].get("content") or ""

    def close(self):
        try:
            self.http.close()
        except Exception:  # noqa: BLE001
            pass


def _prepara(img, lato_max: int):
    from PIL import ImageOps
    img = ImageOps.exif_transpose(img)
    img = img.convert("RGB")
    if max(img.size) > lato_max:
        k = lato_max / max(img.size)
        img = img.resize((max(1, int(img.width * k)), max(1, int(img.height * k))))
    return img


def leggi(path, ocr=None, max_pagine: int = 12, lato_max: int = 1600,
          min_caratteri_pagina: int = 40) -> Estratto:
    """Il testo del file. `ocr`: un oggetto con `leggi(immagine, fonte) -> str` (OcrVisivo, o
    finto nelle prove); senza, le scansioni e le foto danno ErroreTesto."""
    import time
    p = Path(path)
    kind = ESTENSIONI.get(p.suffix.lower())
    if kind is None:
        raise ErroreTesto(f"formato non gestito ({p.suffix})")
    if kind == "testo":
        return Estratto(p.read_text(encoding="utf-8", errors="replace"), "testo")
    if kind == "docx":
        try:
            import docx
        except ImportError:
            raise ErroreTesto("manca python-docx per leggere i file Word") from None
        try:
            d = docx.Document(str(p))
        except Exception as e:  # noqa: BLE001
            raise ErroreTesto(f"file Word rovinato ({type(e).__name__})") from None
        righe = [para.text for para in d.paragraphs if para.text.strip()]
        for t in d.tables:
            for row in t.rows:
                celle = [c.text.strip() for c in row.cells if c.text.strip()]
                if celle:
                    righe.append(" ".join(dict.fromkeys(celle)))
        return Estratto("\n".join(righe), "docx")
    if kind == "immagine":
        try:
            from PIL import Image
        except ImportError:
            raise ErroreTesto("manca Pillow per leggere le immagini") from None
        if ocr is None:
            raise ErroreTesto("è una foto: serve l'OCR (il modello visivo), che qui è spento")
        try:
            im = Image.open(str(p))
        except Exception as e:  # noqa: BLE001
            raise ErroreTesto(f"immagine rovinata ({type(e).__name__})") from None
        testi, n, t0 = [], 0, time.perf_counter()
        for i in range(min(getattr(im, "n_frames", 1), max_pagine)):
            im.seek(i)
            testi.append(ocr.leggi(_prepara(im.copy(), lato_max),
                                   f"{p.name}, pagina {i + 1}").strip())
            n += 1
        return Estratto("\n\n".join(t for t in testi if t), "ocr", n, n,
                        round(time.perf_counter() - t0, 2))
    try:
        import pypdfium2 as pdfium
    except ImportError:
        raise ErroreTesto("manca pypdfium2 per leggere i PDF") from None
    try:
        pdf = pdfium.PdfDocument(str(p))
    except Exception as e:  # noqa: BLE001
        raise ErroreTesto(f"PDF rovinato o protetto ({type(e).__name__})") from None
    try:
        testi, n_ocr, t_ocr, avvisi = [], 0, 0.0, []
        tot = len(pdf)
        for i in range(min(tot, max_pagine)):
            page = pdf[i]
            try:
                txt = page.get_textpage().get_text_range() or ""
            except Exception:  # noqa: BLE001
                txt = ""
            if len(txt.strip()) >= min_caratteri_pagina:
                testi.append(txt.replace("\r\n", "\n").strip())
                continue
            if ocr is None:
                avvisi.append(f"pagina {i + 1} senza testo: serve l'OCR")
                continue
            w = page.get_width() or 595
            img = page.render(scale=lato_max / max(w, page.get_height() or 842)).to_pil()
            t0 = time.perf_counter()
            testi.append(ocr.leggi(_prepara(img, lato_max), f"{p.name}, pagina {i + 1}").strip())
            t_ocr += time.perf_counter() - t0
            n_ocr += 1
        if tot > max_pagine:
            avvisi.append(f"lette le prime {max_pagine} pagine su {tot}")
        if not any(testi) and avvisi and ocr is None:
            raise ErroreTesto("è un PDF scansionato: serve l'OCR (il modello visivo), che qui "
                              "è spento")
        metodo = "pdf+ocr" if n_ocr and len(testi) > n_ocr else ("ocr" if n_ocr else "pdf")
        return Estratto("\n\n".join(t for t in testi if t), metodo, tot, n_ocr,
                        round(t_ocr, 2), avvisi)
    finally:
        pdf.close()
